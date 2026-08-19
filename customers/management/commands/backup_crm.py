from __future__ import annotations

import os
import sqlite3
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand
from django.db import connection
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from customers.models import Customer, FollowUpTask, Project, ProjectProgressUpdate, TomorrowItem, VisitorRecord
from customers.services import FOLLOW_UP_CUSTOMER_GRADES


def display_user(user) -> str:
    if not user:
        return ''
    profile = getattr(user, 'employee_profile', None)
    if profile and profile.nickname:
        return profile.nickname
    return user.get_full_name() or user.username


def excel_value(value):
    if value is None:
        return ''
    if hasattr(value, 'tzinfo') and value.tzinfo is not None:
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def add_sheet(workbook, title, headers, rows, widths=None):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    header_fill = PatternFill('solid', fgColor='176B55')
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for row in rows:
        sheet.append([excel_value(value) for value in row])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    if widths:
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    return sheet


class Command(BaseCommand):
    help = 'Backup the CRM SQLite database and export a readable Excel workbook.'

    def handle(self, *args, **options):
        lock_path = Path(settings.DATA_DIR) / 'backup.lock'
        lock_handle = None
        for attempt in range(2):
            try:
                lock_handle = os.open(
                    lock_path,
                    os.O_CREAT | os.O_EXCL | os.O_WRONLY,
                )
                break
            except FileExistsError:
                is_stale = (
                    lock_path.exists()
                    and timezone.now().timestamp() - lock_path.stat().st_mtime > 3600
                )
                if attempt == 0 and is_stale:
                    lock_path.unlink(missing_ok=True)
                    continue
                return
        try:
            self._perform_backup()
        finally:
            if lock_handle is not None:
                os.close(lock_handle)
            lock_path.unlink(missing_ok=True)

    def _perform_backup(self):
        base_dir = Path(settings.BASE_DIR)
        backup_dir = base_dir / 'backups'
        export_dir = base_dir / 'exports'
        backup_dir.mkdir(parents=True, exist_ok=True)
        export_dir.mkdir(parents=True, exist_ok=True)

        stamp = timezone.localtime().strftime('%Y%m%d-%H%M%S')
        source_database = Path(settings.DATABASES['default']['NAME'])
        backup_database = backup_dir / f'crm-{stamp}.sqlite3'
        with sqlite3.connect(source_database) as source:
            with sqlite3.connect(backup_database) as target:
                source.backup(target)

        workbook = Workbook()
        workbook.remove(workbook.active)

        customers = Customer.objects.select_related(
            'business_owner__employee_profile',
            'technical_owner__employee_profile',
        ).all()
        add_sheet(
            workbook,
            '客资信息',
            ['编号', '等级', '客户名称', '联系电话', '来源', '介绍人',
             '省', '市', '区县', '项目进度', '商务负责人', '技术负责人',
             '客户描述', '施工方案', '创建时间', '更新时间'],
            ([
                customer.id, customer.grade, customer.name, customer.phone,
                customer.source, customer.referrer,
                customer.province, customer.city, customer.district,
                customer.progress,
                display_user(customer.business_owner) or customer.business_owner_name,
                display_user(customer.technical_owner) or customer.technical_owner_name,
                customer.description, customer.plan,
                customer.created_at, customer.updated_at,
            ] for customer in customers),
            [8, 8, 22, 18, 14, 16, 12, 12, 14, 24, 16, 16, 42, 42, 20, 20],
        )

        projects = Project.objects.select_related(
            'customer', 'project_type', 'business_owner__employee_profile',
            'technical_owner__employee_profile',
        ).all()
        add_sheet(
            workbook,
            '项目信息',
            ['编号', '客户', '项目名称', '项目类型', '当前进度', '省', '市', '区县',
             '商务负责人', '技术负责人', '是否有效', '施工方案', '创建时间', '更新时间'],
            ([
                project.id, project.customer.name, project.name,
                project.project_type.name if project.project_type else '未分类', project.progress,
                project.province, project.city, project.district,
                display_user(project.business_owner), display_user(project.technical_owner),
                '是' if project.is_active else '否', project.plan,
                project.created_at, project.updated_at,
            ] for project in projects),
            [8, 22, 30, 14, 24, 12, 12, 14, 16, 16, 10, 48, 20, 20],
        )

        progress_updates = ProjectProgressUpdate.objects.select_related(
            'project__customer', 'project__project_type', 'created_by__employee_profile',
        ).all()
        add_sheet(
            workbook,
            '项目进度更新',
            ['编号', '客户', '项目', '项目类型', '进度内容', '更新人', '更新时间'],
            ([
                item.id, item.project.customer.name, item.project.name,
                item.project.project_type.name if item.project.project_type else '未分类',
                item.content, display_user(item.created_by), item.created_at,
            ] for item in progress_updates),
            [8, 22, 30, 14, 55, 16, 20],
        )

        if VisitorRecord._meta.db_table in connection.introspection.table_names():
            visitor_records = VisitorRecord.objects.select_related(
                'customer', 'host__employee_profile', 'created_by__employee_profile',
            ).all()
            add_sheet(
                workbook,
                '客户来访记录',
                ['编号', '来访日期', '客户', '客户等级', '来访联系人', '来访人数',
                 '来访目的', '接待负责人', '接待记录', '记录人', '创建时间', '更新时间'],
                ([
                    item.id, item.visit_date, item.customer.name, item.customer.grade,
                    item.contact_name, item.visitor_count, item.purpose,
                    display_user(item.host), item.notes, display_user(item.created_by),
                    item.created_at, item.updated_at,
                ] for item in visitor_records),
                [8, 14, 22, 10, 16, 10, 32, 16, 48, 16, 20, 20],
            )

        followups = FollowUpTask.objects.filter(
            project__customer__grade__in=FOLLOW_UP_CUSTOMER_GRADES,
        ).select_related(
            'project__customer', 'assignee__employee_profile',
        ).all()
        add_sheet(
            workbook,
            '跟进记录',
            ['编号', '客户', '项目', '任务内容', '部门', '负责人', '状态',
             '目标阶段', '临时任务', '计划完成', '实际完成', '跟进结果',
             '创建时间', '更新时间'],
            ([
                task.id, task.project.customer.name, task.project.name, task.title,
                task.get_role_display(), display_user(task.assignee),
                task.get_status_display(), task.target_progress,
                '是' if task.is_manual else '否', task.due_at, task.completed_at,
                task.result, task.created_at, task.updated_at,
            ] for task in followups),
            [8, 22, 30, 40, 12, 16, 12, 24, 10, 20, 20, 45, 20, 20],
        )

        tomorrow_items = TomorrowItem.objects.select_related(
            'user__employee_profile',
        ).all()
        add_sheet(
            workbook,
            '明日安排',
            ['编号', '员工', '计划日期', '事项', '补充说明', '是否完成',
             '完成时间', '创建时间', '更新时间'],
            ([
                item.id, display_user(item.user), item.planned_date,
                item.title, item.note, '是' if item.is_completed else '否',
                item.completed_at, item.created_at, item.updated_at,
            ] for item in tomorrow_items),
            [8, 16, 14, 40, 45, 10, 20, 20, 20],
        )

        latest_export = export_dir / '客资总表.xlsx'
        temporary_export = export_dir / f'.客资总表.{os.getpid()}.tmp.xlsx'
        workbook.save(temporary_export)
        try:
            os.replace(temporary_export, latest_export)
        except PermissionError:
            fallback_export = export_dir / f'客资总表-{stamp}.xlsx'
            os.replace(temporary_export, fallback_export)

        backups = sorted(
            backup_dir.glob('crm-*.sqlite3'),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        for old_backup in backups[60:]:
            old_backup.unlink()

        if getattr(self.stdout, '_out', None) is not None:
            self.stdout.write(str(backup_database))
            self.stdout.write(str(latest_export))
