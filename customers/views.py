import json
import math
import re
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from django.utils.http import content_disposition_header
from django.views.decorators.http import require_GET, require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import (
    Customer,
    EmployeeProfile,
    FollowUpTask,
    MaterialExperiment,
    ProgressUpdateReadReceipt,
    ProgressUpdateReadState,
    Project,
    ProjectProgressUpdate,
    ProjectType,
    SchemeCalculation,
    TomorrowItem,
    VisitorRecord,
)
from .services import (
    FOLLOW_UP_CUSTOMER_GRADES,
    PROGRESS_STAGES,
    follow_up_task_matches_project_progress,
    normalize_progress,
    sync_customer_primary_project,
)


@require_GET
def health_check(request):
    try:
        Customer.objects.only('id').exists()
    except Exception:
        return JsonResponse({'status': 'unavailable'}, status=503)
    return JsonResponse({'status': 'ok'})


@require_GET
@login_required
def available_employees(request):
    users = (
        get_user_model()
        .objects.filter(is_active=True)
        .select_related('employee_profile')
        .order_by('first_name', 'username')
    )
    result = {'all': [], 'business': [], 'technical': []}
    for user in users:
        profile = getattr(user, 'employee_profile', None)
        if not user.is_superuser and profile is None:
            continue
        role = EmployeeProfile.Role.BOTH if user.is_superuser else profile.role
        display_name = (
            profile.nickname if profile and profile.nickname
            else user.get_full_name() or user.username
        )
        employee = {
            'id': user.id,
            'username': user.username,
            'display_name': display_name,
            'role': role,
            'businessLevel': (
                profile.business_level
                if profile and role in (EmployeeProfile.Role.BUSINESS, EmployeeProfile.Role.BOTH)
                else (5 if user.is_superuser else None)
            ),
            'technicalLevel': (
                profile.technical_level
                if profile and role in (EmployeeProfile.Role.TECHNICAL, EmployeeProfile.Role.BOTH)
                else (5 if user.is_superuser else None)
            ),
        }
        result['all'].append(employee)
        if role in (EmployeeProfile.Role.BUSINESS, EmployeeProfile.Role.BOTH):
            result['business'].append(employee)
        if role in (EmployeeProfile.Role.TECHNICAL, EmployeeProfile.Role.BOTH):
            result['technical'].append(employee)
    return JsonResponse(result)


def employee_display_name(user, fallback='待分配'):
    if not user:
        return fallback
    profile = getattr(user, 'employee_profile', None)
    return (
        profile.nickname if profile and profile.nickname
        else user.get_full_name() or user.username
    )


def can_delete_customer(user):
    if user.is_superuser:
        return True
    profile = getattr(user, 'employee_profile', None)
    return bool(
        profile
        and profile.role in (
            EmployeeProfile.Role.BUSINESS,
            EmployeeProfile.Role.BOTH,
        )
    )


def can_manage_project(user, project):
    return (
        user.is_superuser
        or project.business_owner_id == user.id
        or project.technical_owner_id == user.id
    )


def can_manage_follow_up(user, task):
    return (
        user.is_superuser
        or task.assignee_id == user.id
        or can_manage_project(user, task.project)
    )


def parse_task_due_at(value):
    raw_value = str(value or '').strip()
    if not raw_value:
        return None
    due_at = parse_datetime(raw_value)
    if due_at is None:
        return None
    if timezone.is_naive(due_at):
        due_at = timezone.make_aware(due_at, timezone.get_current_timezone())
    return due_at


def stage_percent(progress):
    normalized = normalize_progress(progress)
    try:
        index = PROGRESS_STAGES.index(normalized)
    except ValueError:
        return 5
    return round(5 + index / (len(PROGRESS_STAGES) - 1) * 95)


def serialize_customer(customer):
    created_at = timezone.localtime(customer.created_at)
    owner_name = employee_display_name(
        customer.business_owner,
        customer.business_owner_name or '待分配',
    )
    tech_name = employee_display_name(
        customer.technical_owner,
        customer.technical_owner_name or '待分配',
    )
    primary_project = customer.projects.select_related('project_type').order_by('created_at', 'id').first()
    return {
        'id': customer.id,
        'name': customer.name,
        'phone': customer.phone,
        'source': customer.source,
        'channel': customer.channel,
        'referrer': customer.referrer,
        'province': customer.province,
        'city': customer.city,
        'district': customer.district,
        'region': ' · '.join(filter(None, (customer.province, customer.city))),
        'grade': customer.grade,
        'description': customer.description,
        'progress': customer.progress,
        'percent': stage_percent(customer.progress),
        'plan': customer.plan,
        'projectTypeId': primary_project.project_type_id if primary_project else '',
        'projectTypeName': primary_project.project_type.name if primary_project and primary_project.project_type else '未分类',
        'ownerId': customer.business_owner_id or '',
        'owner': owner_name,
        'techId': customer.technical_owner_id or '',
        'tech': tech_name,
        'createdAt': customer.created_at.isoformat(),
        'createdDate': created_at.strftime('%Y-%m-%d'),
        'created': created_at.strftime('%Y年%m月%d日'),
        'updated': customer.updated_at.strftime('%m月%d日 %H:%M'),
        'color': customer.color or '#5b7cfa',
    }


def parse_payload(request):
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return None, JsonResponse({'error': '请求数据格式不正确'}, status=400)
    if not isinstance(payload, dict):
        return None, JsonResponse({'error': '请求数据格式不正确'}, status=400)
    return payload, None


def get_available_owner(user_id, allowed_roles):
    if not user_id:
        return None
    user = (
        get_user_model()
        .objects.filter(id=user_id, is_active=True)
        .select_related('employee_profile')
        .first()
    )
    if not user:
        return None
    if user.is_superuser:
        return user
    profile = getattr(user, 'employee_profile', None)
    if not profile or profile.role not in allowed_roles:
        return None
    return user


def update_customer_from_payload(customer, payload):
    required = {
        'name': '客户名称',
        'province': '省份',
        'city': '城市',
    }
    missing = [
        label for field, label in required.items()
        if not str(payload.get(field, '')).strip()
    ]
    if missing:
        return f"请填写：{'、'.join(missing)}"

    source = str(payload.get('source', '抖音')).strip()
    referrer = str(payload.get('referrer', '')).strip() if source == '朋友介绍' else ''
    if source == '朋友介绍' and not referrer:
        return '朋友介绍的客资必须填写介绍人'

    owner_id = payload.get('ownerId')
    technician_id = payload.get('techId')
    owner = None
    technician = None
    if owner_id:
        owner = get_available_owner(
            owner_id,
            (EmployeeProfile.Role.BUSINESS, EmployeeProfile.Role.BOTH),
        )
        if not owner:
            return '所选商务负责人账号不可用'
    if technician_id:
        technician = get_available_owner(
            technician_id,
            (EmployeeProfile.Role.TECHNICAL, EmployeeProfile.Role.BOTH),
        )
        if not technician:
            return '所选技术负责人账号不可用'

    customer.name = str(payload.get('name', '')).strip()
    customer.phone = str(payload.get('phone', '')).strip()
    customer.source = source
    customer.referrer = referrer
    customer.channel = f'介绍人：{referrer}' if referrer else '手动修改'
    customer.province = str(payload.get('province', '')).strip()
    customer.city = str(payload.get('city', '')).strip()
    customer.district = str(payload.get('district', '')).strip()
    grade = str(payload.get('grade', 'D')).strip().upper()
    customer.grade = grade if grade in ('A', 'B', 'C', 'D') else 'D'
    customer.description = str(payload.get('description', '')).strip()
    progress = str(payload.get('progress', '需求对接')).strip()
    if progress not in PROGRESS_STAGES:
        return '请选择有效的项目进度'
    customer.progress = progress
    customer.plan = str(payload.get('plan', '')).strip() or '方案待完善'
    if owner:
        customer.business_owner = owner
        customer.business_owner_name = employee_display_name(owner)
    if technician:
        customer.technical_owner = technician
        customer.technical_owner_name = employee_display_name(technician)
    return ''


def project_type_from_payload(payload, *, required_key=False):
    if 'projectTypeId' not in payload:
        return ('preserve' if not required_key else None), ''
    project_type_id = payload.get('projectTypeId')
    if project_type_id in ('', None):
        return None, ''
    project_type = ProjectType.objects.filter(id=project_type_id, is_active=True).first()
    if not project_type:
        return None, '请选择有效的项目类型'
    return project_type.id, ''


@login_required
@require_http_methods(['GET', 'POST'])
def customers_collection(request):
    if request.method == 'GET':
        customers = Customer.objects.select_related(
            'business_owner__employee_profile',
            'technical_owner__employee_profile',
        ).prefetch_related('projects__project_type')
        projects = Project.objects.select_related(
            'customer',
            'project_type',
            'business_owner__employee_profile',
            'technical_owner__employee_profile',
        )
        return JsonResponse({
            'customers': [serialize_customer(item) for item in customers],
            'projects': [serialize_project(item) for item in projects],
        })

    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    customer = Customer(sort_order=Customer.objects.count())
    validation_error = update_customer_from_payload(customer, payload)
    if validation_error:
        return JsonResponse({'error': validation_error}, status=400)
    project_type_id, validation_error = project_type_from_payload(payload, required_key=True)
    if validation_error:
        return JsonResponse({'error': validation_error}, status=400)
    customer.save()
    sync_customer_primary_project(customer, project_type_id)
    return JsonResponse({'customer': serialize_customer(customer)}, status=201)


@login_required
@require_http_methods(['POST'])
def customers_import(request):
    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    items = payload.get('customers')
    if not isinstance(items, list) or not items:
        return JsonResponse({'error': '没有找到可导入的有效客资'}, status=400)
    if len(items) > 500:
        return JsonResponse({'error': '单次最多导入500条客资'}, status=400)

    created_customers = []
    try:
        with transaction.atomic():
            start_order = Customer.objects.count()
            for index, item in enumerate(items):
                row_number = index + 2
                if not isinstance(item, dict):
                    raise ValueError(f'第{row_number}行数据格式不正确')
                customer = Customer(sort_order=start_order + index)
                validation_error = update_customer_from_payload(customer, item)
                if validation_error:
                    raise ValueError(f'第{row_number}行：{validation_error}')
                project_type_id, validation_error = project_type_from_payload(
                    item,
                    required_key=True,
                )
                if validation_error:
                    raise ValueError(f'第{row_number}行：{validation_error}')
                customer.save()
                sync_customer_primary_project(customer, project_type_id)
                created_customers.append(customer)
    except ValueError as error:
        return JsonResponse({'error': str(error)}, status=400)

    return JsonResponse({
        'createdCount': len(created_customers),
        'customers': [serialize_customer(customer) for customer in created_customers],
    }, status=201)


@login_required
@require_http_methods(['GET', 'PATCH', 'DELETE'])
def customer_detail(request, customer_id):
    customer = get_object_or_404(
        Customer.objects.select_related(
            'business_owner__employee_profile',
            'technical_owner__employee_profile',
        ),
        id=customer_id,
    )
    if request.method == 'DELETE':
        if not can_delete_customer(request.user):
            return JsonResponse({'error': '只有商务部人员或超级管理员可以删除客资'}, status=403)
        customer_name = customer.name
        project_count = customer.projects.count()
        follow_up_count = FollowUpTask.objects.filter(project__customer=customer).count()
        customer.delete()
        return JsonResponse({
            'deleted': True,
            'customerId': customer_id,
            'customerName': customer_name,
            'projectCount': project_count,
            'followUpCount': follow_up_count,
        })
    if request.method == 'GET':
        follow_ups = FollowUpTask.objects.none()
        progress_update_items = ProjectProgressUpdate.objects.none()
        project = primary_customer_project(customer)
        if customer.grade in FOLLOW_UP_CUSTOMER_GRADES:
            follow_up_items = FollowUpTask.objects.filter(
                project__customer_id=customer.id,
            ).select_related(
                'assignee__employee_profile',
                'project__customer',
                'project__business_owner__employee_profile',
                'project__technical_owner__employee_profile',
            ).order_by('-completed_at', '-updated_at')
            follow_ups = [
                task for task in follow_up_items
                if follow_up_task_matches_project_progress(task)
            ]
            progress_update_items = ProjectProgressUpdate.objects.filter(
                project__customer_id=customer.id,
                project__is_active=True,
            ).select_related(
                'created_by__employee_profile',
                'project__customer',
                'project__project_type',
                'project__business_owner__employee_profile',
                'project__technical_owner__employee_profile',
            ).order_by('occurred_at', 'id')
        visitor_items = visitor_record_queryset().filter(customer_id=customer.id)
        scheme_item = None
        experiment_item = None
        if project:
            scheme_item = SchemeCalculation.objects.select_related(
                'created_by__employee_profile',
                'updated_by__employee_profile',
            ).filter(project=project).first()
            experiment_item = MaterialExperiment.objects.select_related(
                'created_by__employee_profile',
                'updated_by__employee_profile',
            ).filter(project=project).first()
        return JsonResponse({
            'customer': serialize_customer(customer),
            'followUps': [serialize_follow_up_task(task, request.user) for task in follow_ups],
            'progressUpdates': [serialize_progress_update(item, request.user) for item in progress_update_items],
            'visitorRecords': [serialize_visitor_record(item, request.user) for item in visitor_items],
            'schemeCalculation': (
                serialize_scheme_calculation(scheme_item, request.user)
                if scheme_item else None
            ),
            'materialExperiment': (
                serialize_material_experiment(experiment_item, project, request.user)
                if project else None
            ),
            'canManageProject': bool(project and can_manage_project(request.user, project)),
        })

    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    validation_error = update_customer_from_payload(customer, payload)
    if validation_error:
        return JsonResponse({'error': validation_error}, status=400)
    project_type_id, validation_error = project_type_from_payload(payload)
    if validation_error:
        return JsonResponse({'error': validation_error}, status=400)
    customer.save()
    sync_customer_primary_project(customer, project_type_id)
    customer = Customer.objects.select_related(
        'business_owner__employee_profile',
        'technical_owner__employee_profile',
    ).get(id=customer.id)
    return JsonResponse({'customer': serialize_customer(customer)})


def primary_customer_project(customer):
    return customer.projects.select_related(
        'customer',
        'project_type',
        'business_owner__employee_profile',
        'technical_owner__employee_profile',
    ).order_by('created_at', 'id').first()


def decimal_value(value, label, *, minimum=Decimal('0.0001'), allow_zero=False):
    try:
        result = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(f'{label}必须填写有效数字')
    if allow_zero and result == 0:
        return result
    if result < minimum:
        raise ValueError(f'{label}必须大于0')
    return result


def normalized_mix_description(description, dosage, material_name):
    """Keep the dosage shown in the mix description aligned with its numeric field."""
    dosage_text = f'{float(dosage):g}%'
    base_description = re.sub(
        r'^\s*\d+(?:\.\d+)?\s*%\s*',
        '',
        str(description or '').strip(),
    )
    if not base_description:
        base_description = f'{material_name}（PS-I）'
    return f'{dosage_text}{base_description}'[:200]


def normalize_scheme_payload(payload):
    try:
        layer_count = int(payload.get('layerCount', 1))
    except (TypeError, ValueError):
        raise ValueError('请选择1至3层结构')
    if layer_count not in (1, 2, 3):
        raise ValueError('结构层数只能选择1层、2层或3层')
    source_layers = payload.get('layers')
    if not isinstance(source_layers, list) or len(source_layers) < layer_count:
        raise ValueError(f'请完整填写{layer_count}层结构数据')

    layers = []
    for index, raw_layer in enumerate(source_layers[:layer_count], start=1):
        if not isinstance(raw_layer, dict):
            raise ValueError(f'第{index}层数据格式不正确')
        dosage = decimal_value(raw_layer.get('dosagePercent'), f'第{index}层掺量')
        if dosage > 100:
            raise ValueError(f'第{index}层掺量不能超过100%')
        unit_price = decimal_value(
            raw_layer.get('unitPrice'),
            f'第{index}层单价',
            minimum=Decimal('0'),
            allow_zero=True,
        )
        material_name = str(raw_layer.get('name') or '土凝岩稳定土').strip()[:100]
        layers.append({
            'name': material_name,
            'structureLayer': str(raw_layer.get('structureLayer') or f'第{index}层').strip()[:50],
            'mixDescription': normalized_mix_description(
                raw_layer.get('mixDescription'), dosage, material_name,
            ),
            'length': float(decimal_value(raw_layer.get('length'), f'第{index}层长度')),
            'width': float(decimal_value(raw_layer.get('width'), f'第{index}层宽度')),
            'thickness': float(decimal_value(raw_layer.get('thickness'), f'第{index}层厚度')),
            'dosagePercent': float(dosage),
            'unitPrice': float(unit_price),
            'density': float(decimal_value(raw_layer.get('density'), f'第{index}层密度')),
        })

    return {
        'title': str(payload.get('title') or '').strip()[:150],
        'layerCount': layer_count,
        'layers': layers,
        'remarks': str(payload.get('remarks') or '').strip(),
    }


def calculate_scheme_layer(layer):
    dosage_ratio = Decimal(str(layer['dosagePercent'])) / Decimal('100')
    exact_quantity = (
        Decimal(str(layer['length']))
        * Decimal(str(layer['width']))
        * Decimal(str(layer['thickness']))
        * dosage_ratio
        * Decimal(str(layer['density']))
    )
    rounded_quantity = math.ceil(exact_quantity)
    square_price = (
        Decimal(str(layer['thickness']))
        * dosage_ratio
        * Decimal(str(layer['density']))
        * Decimal(str(layer['unitPrice']))
    )
    total_price = Decimal(rounded_quantity) * Decimal(str(layer['unitPrice']))
    return {
        **layer,
        'exactQuantity': round(float(exact_quantity), 3),
        'quantity': rounded_quantity,
        'squareMeterPrice': round(float(square_price), 2),
        'totalPrice': round(float(total_price), 2),
    }


def serialize_scheme_calculation(calculation, user):
    layers = [calculate_scheme_layer(layer) for layer in calculation.layers]
    return {
        'id': calculation.id,
        'projectId': calculation.project_id,
        'title': calculation.title,
        'layerCount': calculation.layer_count,
        'layers': layers,
        'remarks': calculation.remarks,
        'totalPrice': round(sum(layer['totalPrice'] for layer in layers), 2),
        'createdBy': employee_display_name(calculation.created_by, '系统'),
        'updatedBy': employee_display_name(calculation.updated_by, '系统'),
        'createdLabel': timezone.localtime(calculation.created_at).strftime('%Y年%m月%d日 %H:%M'),
        'updatedLabel': timezone.localtime(calculation.updated_at).strftime('%Y年%m月%d日 %H:%M'),
        'canEdit': can_manage_project(user, calculation.project),
    }


def serialize_material_experiment(experiment, project, user):
    return {
        'id': experiment.id if experiment else None,
        'projectId': project.id,
        'stableMaterial': experiment.stable_material if experiment else '',
        'day7Data': experiment.day_7_data if experiment else '',
        'day14Data': experiment.day_14_data if experiment else '',
        'day28Data': experiment.day_28_data if experiment else '',
        'technicalMessage': experiment.technical_message if experiment else '',
        'createdBy': employee_display_name(experiment.created_by, '系统') if experiment else '',
        'updatedBy': employee_display_name(experiment.updated_by, '系统') if experiment else '',
        'createdLabel': (
            timezone.localtime(experiment.created_at).strftime('%Y年%m月%d日 %H:%M')
            if experiment else ''
        ),
        'updatedLabel': (
            timezone.localtime(experiment.updated_at).strftime('%Y年%m月%d日 %H:%M')
            if experiment else ''
        ),
        'canEdit': can_manage_project(user, project),
    }


def default_scheme_payload(project):
    return {
        'id': None,
        'projectId': project.id,
        'title': f'{project.customer.name}{project.project_type.name if project.project_type else "项目"}造价分析',
        'layerCount': 1,
        'layers': [{
            'name': '土凝岩稳定土',
            'structureLayer': '面层',
            'mixDescription': '土凝岩稳定土（PS-I）',
            'length': '',
            'width': '',
            'thickness': 0.2,
            'dosagePercent': 10,
            'unitPrice': 600,
            'density': 1.7,
            'exactQuantity': 0,
            'quantity': 0,
            'squareMeterPrice': 20.4,
            'totalPrice': 0,
        }],
        'remarks': '实际掺配比例需根据土质情况结合试验确定。\n此次报价不含运输费用。\n土凝岩材料报价包含13%税点。',
        'totalPrice': 0,
        'createdBy': '',
        'updatedBy': '',
        'createdLabel': '',
        'updatedLabel': '',
        'canEdit': True,
    }


@login_required
@require_http_methods(['GET', 'PUT'])
def scheme_calculation(request, customer_id):
    customer = get_object_or_404(Customer, id=customer_id)
    project = primary_customer_project(customer)
    if not project:
        return JsonResponse({'error': '该客户尚未建立项目，无法进行方案测算'}, status=400)

    if request.method == 'GET':
        calculation = SchemeCalculation.objects.select_related(
            'project__customer',
            'project__project_type',
            'project__business_owner',
            'project__technical_owner',
            'created_by__employee_profile',
            'updated_by__employee_profile',
        ).filter(project=project).first()
        if not calculation:
            data = default_scheme_payload(project)
            data['canEdit'] = can_manage_project(request.user, project)
            return JsonResponse({'calculation': data})
        return JsonResponse({'calculation': serialize_scheme_calculation(calculation, request.user)})

    if not can_manage_project(request.user, project):
        return JsonResponse({'error': '只有该项目负责人或超级管理员可以保存方案测算'}, status=403)
    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    try:
        normalized = normalize_scheme_payload(payload)
    except ValueError as error:
        return JsonResponse({'error': str(error)}, status=400)
    normalized['title'] = normalized['title'] or f'{customer.name}项目造价分析'
    calculation, created = SchemeCalculation.objects.get_or_create(
        project=project,
        defaults={
            'title': normalized['title'],
            'layer_count': normalized['layerCount'],
            'layers': normalized['layers'],
            'remarks': normalized['remarks'],
            'created_by': request.user,
            'updated_by': request.user,
        },
    )
    if not created:
        calculation.title = normalized['title']
        calculation.layer_count = normalized['layerCount']
        calculation.layers = normalized['layers']
        calculation.remarks = normalized['remarks']
        calculation.updated_by = request.user
        calculation.save()
    return JsonResponse({'calculation': serialize_scheme_calculation(calculation, request.user)})


@login_required
@require_http_methods(['GET', 'PUT'])
def material_experiment(request, customer_id):
    customer = get_object_or_404(Customer, id=customer_id)
    project = primary_customer_project(customer)
    if not project:
        return JsonResponse({'error': '该客户尚未建立项目，无法登记实验数据'}, status=400)
    experiment = MaterialExperiment.objects.select_related(
        'created_by__employee_profile',
        'updated_by__employee_profile',
    ).filter(project=project).first()
    if request.method == 'GET':
        return JsonResponse({
            'experiment': serialize_material_experiment(experiment, project, request.user),
        })
    if not can_manage_project(request.user, project):
        return JsonResponse({'error': '只有该项目负责人或超级管理员可以保存实验数据'}, status=403)
    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    values = {
        'stable_material': str(payload.get('stableMaterial', '')).strip()[:150],
        'day_7_data': str(payload.get('day7Data', '')).strip()[:2000],
        'day_14_data': str(payload.get('day14Data', '')).strip()[:2000],
        'day_28_data': str(payload.get('day28Data', '')).strip()[:2000],
        'technical_message': str(payload.get('technicalMessage', '')).strip()[:3000],
    }
    if not any(values.values()):
        return JsonResponse({'error': '请至少填写一项实验数据'}, status=400)
    if experiment is None:
        experiment = MaterialExperiment.objects.create(
            project=project,
            created_by=request.user,
            updated_by=request.user,
            **values,
        )
    else:
        for field, value in values.items():
            setattr(experiment, field, value)
        experiment.updated_by = request.user
        experiment.save()
    return JsonResponse({
        'experiment': serialize_material_experiment(experiment, project, request.user),
    })


def scheme_filename(title):
    safe_title = re.sub(r'[\\/:*?"<>|\r\n]+', '_', title).strip(' ._') or '施工方案测算'
    return f'{safe_title}.xlsx'


def build_scheme_workbook(calculation):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '方案测算'
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A13'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35

    headers = ['序号', '名称', '结构层', '配比说明', '长度m', '宽度m', '厚度m', '掺量', '单价（元/吨）', '土凝岩单平米报价（元/㎡）', '土凝岩用量（吨）', '材料总价（元）']
    widths = [7, 20, 11, 26, 11, 11, 11, 11, 15, 22, 17, 19]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    dark_green = '0B3D2E'
    brand_green = '0F7658'
    mint = 'E9F4EF'
    pale_mint = 'F5FAF7'
    gold = 'D5A33D'
    pale_gold = 'FFF4D6'
    ink = '193B33'
    muted = '647A73'
    white = 'FFFFFF'
    line_color = 'BDD5CC'
    border = Border(*(Side(style='thin', color=line_color) for _ in range(4)))

    sheet.merge_cells('A1:L2')
    sheet['A1'] = calculation.title
    sheet['A1'].font = Font(name='微软雅黑', size=21, bold=True, color=white)
    sheet['A1'].fill = PatternFill('solid', fgColor=dark_green)
    sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 30

    sheet.merge_cells('A3:L3')
    sheet['A3'] = '银鼎土凝岩  ·  项目材料方案报价'
    sheet['A3'].font = Font(name='微软雅黑', size=10, bold=True, color=dark_green)
    sheet['A3'].fill = PatternFill('solid', fgColor='DCECE5')
    sheet['A3'].alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[3].height = 24

    project = calculation.project
    customer = project.customer
    region = ' / '.join(filter(None, (project.province or customer.province, project.city or customer.city, project.district or customer.district))) or '待补充'
    customer_phone = customer.phone or '待补充'
    business_owner = employee_display_name(project.business_owner)
    prepared_date = timezone.localdate().strftime('%Y年%m月%d日')
    info_rows = [
        ('客户单位', customer.name, '项目名称', project.name),
        ('项目地区', region, '商务负责人', business_owner),
        ('联系电话', customer_phone, '编制日期', prepared_date),
    ]
    for row, (left_label, left_value, right_label, right_value) in enumerate(info_rows, start=4):
        for start_column, end_column, value, is_label in (
            (1, 2, left_label, True), (3, 6, left_value, False),
            (7, 8, right_label, True), (9, 12, right_value, False),
        ):
            sheet.merge_cells(start_row=row, start_column=start_column, end_row=row, end_column=end_column)
            cell = sheet.cell(row=row, column=start_column, value=value)
            cell.font = Font(name='微软雅黑', size=10, bold=is_label, color=ink if is_label else '233F38')
            cell.fill = PatternFill('solid', fgColor=mint if is_label else white)
            cell.alignment = Alignment(horizontal='left', vertical='center')
            for column in range(start_column, end_column + 1):
                sheet.cell(row=row, column=column).border = border
        sheet.row_dimensions[row].height = 25

    layers = [calculate_scheme_layer(layer) for layer in calculation.layers[:calculation.layer_count]]
    dosage_text = ' + '.join(f"{layer['dosagePercent']:g}%" for layer in layers) or '—'
    total_square_price = sum(layer['squareMeterPrice'] for layer in layers)
    square_price_text = f'¥{total_square_price:,.2f} /㎡' if layers else '—'
    total_quantity = sum(layer['quantity'] for layer in layers)
    total_price = sum(layer['totalPrice'] for layer in layers)
    summary_cards = [
        (1, 3, '推荐掺量', dosage_text, pale_gold, gold, dark_green),
        (4, 6, '各层单平米报价合计', square_price_text, pale_gold, gold, dark_green),
        (7, 9, '土凝岩总用量', f'{total_quantity:,.0f} 吨', mint, brand_green, dark_green),
        (10, 12, '材料总价', f'¥{total_price:,.2f}', dark_green, gold, 'FF5A5F'),
    ]
    for start_column, end_column, label, value, fill_color, accent_color, value_color in summary_cards:
        sheet.merge_cells(start_row=8, start_column=start_column, end_row=8, end_column=end_column)
        sheet.merge_cells(start_row=9, start_column=start_column, end_row=10, end_column=end_column)
        label_cell = sheet.cell(row=8, column=start_column, value=label)
        value_cell = sheet.cell(row=9, column=start_column, value=value)
        label_cell.font = Font(name='微软雅黑', size=9, bold=True, color=accent_color)
        value_cell.font = Font(name='微软雅黑', size=16, bold=True, color=value_color)
        label_cell.alignment = Alignment(horizontal='center', vertical='center')
        value_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in range(8, 11):
            for column in range(start_column, end_column + 1):
                cell = sheet.cell(row=row, column=column)
                cell.fill = PatternFill('solid', fgColor=fill_color)
                cell.border = Border(
                    left=Side(style='medium' if column == start_column else 'thin', color=accent_color if column == start_column else line_color),
                    right=Side(style='thin', color=line_color),
                    top=Side(style='thin', color=line_color),
                    bottom=Side(style='thin', color=line_color),
                )
    sheet.row_dimensions[8].height = 22
    sheet.row_dimensions[9].height = 25
    sheet.row_dimensions[10].height = 25

    header_row = 12
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(name='微软雅黑', size=9, bold=True, color=white)
        cell.fill = PatternFill('solid', fgColor=gold if column in (8, 9, 10, 12) else brand_green)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    sheet.row_dimensions[header_row].height = 38

    first_row = header_row + 1
    for offset, layer in enumerate(layers):
        row = first_row + offset
        mix_description = normalized_mix_description(
            layer.get('mixDescription'), layer['dosagePercent'], layer['name'],
        )
        values = [
            offset + 1,
            layer['name'],
            layer['structureLayer'],
            mix_description,
            layer['length'],
            layer['width'],
            layer['thickness'],
            layer['dosagePercent'] / 100,
            layer['unitPrice'],
            layer['squareMeterPrice'],
            layer['quantity'],
            layer['totalPrice'],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name='微软雅黑', size=10, bold=column in (8, 9, 10, 11, 12), color=dark_green if column != 12 else white)
            if column in (8, 9, 10):
                cell.fill = PatternFill('solid', fgColor=pale_gold)
            elif column == 11:
                cell.fill = PatternFill('solid', fgColor=mint)
            elif column == 12:
                cell.fill = PatternFill('solid', fgColor=brand_green)
            else:
                cell.fill = PatternFill('solid', fgColor=white if offset % 2 == 0 else pale_mint)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        sheet.cell(row=row, column=8).number_format = '0.00%'
        for column in (5, 6, 7, 9, 10, 12):
            sheet.cell(row=row, column=column).number_format = '#,##0.00'
        sheet.cell(row=row, column=11).number_format = '0'
        sheet.cell(row=row, column=12).number_format = '¥#,##0.00'
        sheet.row_dimensions[row].height = 44

    total_row = first_row + len(layers)
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
    sheet.cell(row=total_row, column=1, value='报价合计')
    sheet.cell(row=total_row, column=8, value=dosage_text)
    sheet.merge_cells(start_row=total_row, start_column=9, end_row=total_row, end_column=10)
    sheet.cell(row=total_row, column=9, value='土凝岩用量')
    sheet.cell(row=total_row, column=11, value=total_quantity)
    sheet.cell(row=total_row, column=12, value=total_price)
    for column in range(1, 13):
        cell = sheet.cell(row=total_row, column=column)
        cell.font = Font(name='微软雅黑', size=11, bold=True, color='FF5A5F' if column == 12 else (white if column == 1 else dark_green))
        cell.fill = PatternFill('solid', fgColor=dark_green if column in range(1, 8) or column == 12 else pale_gold)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    sheet.cell(row=total_row, column=8).font = Font(name='微软雅黑', size=12, bold=True, color=dark_green)
    sheet.cell(row=total_row, column=11).number_format = '0 "吨"'
    sheet.cell(row=total_row, column=12).number_format = '¥#,##0.00'
    sheet.row_dimensions[total_row].height = 34

    note_title_row = total_row + 2
    sheet.merge_cells(start_row=note_title_row, start_column=1, end_row=note_title_row, end_column=12)
    note_title = sheet.cell(row=note_title_row, column=1, value='报价备注 / TERMS & NOTES')
    note_title.font = Font(name='微软雅黑', size=10, bold=True, color=white)
    note_title.fill = PatternFill('solid', fgColor=brand_green)
    note_title.alignment = Alignment(horizontal='left', vertical='center')
    note_title.border = border
    sheet.row_dimensions[note_title_row].height = 26

    note_row = note_title_row + 1
    remark_lines = [line.strip() for line in (calculation.remarks or '').splitlines() if line.strip()]
    display_remarks = remark_lines or ['无']
    while len(display_remarks) < 3:
        display_remarks.append('')
    for offset, remark in enumerate(display_remarks):
        row = note_row + offset
        number_cell = sheet.cell(row=row, column=1, value=offset + 1)
        number_cell.font = Font(name='微软雅黑', size=11, bold=True, color=white)
        number_cell.fill = PatternFill('solid', fgColor=gold)
        number_cell.alignment = Alignment(horizontal='center', vertical='center')
        number_cell.border = border
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=12)
        remark_cell = sheet.cell(row=row, column=2, value=remark)
        remark_cell.font = Font(name='微软雅黑', size=10, color=ink)
        remark_cell.fill = PatternFill('solid', fgColor=pale_mint if offset % 2 == 0 else white)
        remark_cell.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True)
        for column in range(2, 13):
            sheet.cell(row=row, column=column).border = border
        sheet.row_dimensions[row].height = 30

    sheet.auto_filter.ref = None
    sheet.print_title_rows = '1:12'
    sheet.print_area = f'A1:L{note_row + len(display_remarks) - 1}'
    sheet.oddFooter.center.text = '银鼎土凝岩 · 专业材料方案'
    sheet.oddFooter.center.size = 8
    sheet.oddFooter.center.color = muted
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook


@login_required
@require_http_methods(['POST'])
def export_scheme_calculation(request, customer_id):
    customer = get_object_or_404(Customer, id=customer_id)
    project = primary_customer_project(customer)
    if not project:
        return JsonResponse({'error': '该客户尚未建立项目'}, status=400)
    if not can_manage_project(request.user, project):
        return JsonResponse({'error': '只有该项目负责人或超级管理员可以导出方案测算'}, status=403)
    calculation = SchemeCalculation.objects.select_related('project__customer').filter(project=project).first()
    if not calculation:
        return JsonResponse({'error': '请先保存方案测算，再执行导出'}, status=400)
    try:
        filename = scheme_filename(calculation.title)
        output = BytesIO()
        build_scheme_workbook(calculation).save(output)
    except OSError as error:
        return JsonResponse({
            'error': f'Excel 生成失败：{error}',
        }, status=500)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = content_disposition_header(True, filename)
    response['Content-Length'] = output.tell()
    return response


def serialize_project(project):
    created_at = timezone.localtime(project.created_at)
    return {
        'id': project.id,
        'name': project.name,
        'progress': project.progress,
        'plan': project.plan,
        'projectType': {
            'id': project.project_type_id or '',
            'name': project.project_type.name if project.project_type else '未分类',
        },
        'isActive': project.is_active,
        'region': ' · '.join(filter(None, (project.province, project.city, project.district))),
        'province': project.province,
        'city': project.city,
        'district': project.district,
        'customer': {
            'id': project.customer_id,
            'name': project.customer.name,
            'grade': project.customer.grade,
            'phone': project.customer.phone,
            'source': project.customer.source,
        },
        'businessOwner': {
            'id': project.business_owner_id or '',
            'name': employee_display_name(project.business_owner),
            'businessLevel': (
                project.business_owner.employee_profile.business_level
                if project.business_owner and hasattr(project.business_owner, 'employee_profile')
                else (5 if project.business_owner and project.business_owner.is_superuser else 1)
            ),
        },
        'technicalOwner': {
            'id': project.technical_owner_id or '',
            'name': employee_display_name(project.technical_owner),
            'technicalLevel': (
                project.technical_owner.employee_profile.technical_level
                if project.technical_owner and hasattr(project.technical_owner, 'employee_profile')
                else (5 if project.technical_owner and project.technical_owner.is_superuser else 1)
            ),
        },
        'createdDate': created_at.strftime('%Y-%m-%d'),
        'created': created_at.strftime('%Y年%m月%d日'),
    }


def serialize_follow_up_task(task, user=None):
    now = timezone.now()
    target_progress = normalize_progress(task.target_progress)
    try:
        target_index = PROGRESS_STAGES.index(target_progress)
    except ValueError:
        target_index = 0
    from_progress = (
        target_progress
        if task.is_manual
        else PROGRESS_STAGES[max(0, target_index - 1)]
    )
    title = (
        '组织项目实施并记录现场进展'
        if not task.is_manual and task.target_progress in ('A段施工完成', 'B段施工完成')
        else task.title
    )
    is_overdue = (
        task.status != FollowUpTask.Status.COMPLETED
        and task.due_at is not None
        and task.due_at < now
    )
    return {
        'id': task.id,
        'title': title,
        'fromProgress': from_progress,
        'targetProgress': target_progress,
        'isManual': task.is_manual,
        'canEdit': bool(user and can_manage_follow_up(user, task)),
        'canDelete': bool(user and can_manage_follow_up(user, task)),
        'canCreateManual': bool(user and can_manage_project(user, task.project)),
        'role': task.role,
        'roleLabel': task.get_role_display(),
        'status': task.status,
        'statusLabel': task.get_status_display(),
        'result': task.result,
        'completedAt': task.completed_at.isoformat() if task.completed_at else '',
        'completedLabel': (
            timezone.localtime(task.completed_at).strftime('%Y年%m月%d日 %H:%M')
            if task.completed_at else ''
        ),
        'updatedLabel': timezone.localtime(task.updated_at).strftime('%Y年%m月%d日 %H:%M'),
        'dueAt': task.due_at.isoformat() if task.due_at else '',
        'dueLabel': timezone.localtime(task.due_at).strftime('%m月%d日 %H:%M') if task.due_at else '未设置',
        'isOverdue': is_overdue,
        'assignee': {
            'id': task.assignee_id or '',
            'name': employee_display_name(
                task.assignee,
                (
                    task.project.customer.business_owner_name
                    if task.role == FollowUpTask.Role.BUSINESS
                    else task.project.customer.technical_owner_name
                ) or '待分配',
            ),
            'username': task.assignee.username if task.assignee else '',
        },
        'project': serialize_project(task.project),
    }


def collapse_tasks_by_project(tasks, now):
    """Return one current status card per project while preserving task history in the database."""
    grouped = {}
    for task in tasks:
        grouped.setdefault(task.project_id, []).append(task)

    collapsed = []
    for project_tasks in grouped.values():
        active_tasks = [
            task for task in project_tasks
            if task.status != FollowUpTask.Status.COMPLETED
        ]
        if active_tasks:
            active_tasks.sort(key=lambda task: (
                0 if task.due_at and task.due_at < now else 1,
                0 if task.status == FollowUpTask.Status.PENDING else 1,
                task.due_at or now + timedelta(days=36500),
                -task.updated_at.timestamp(),
            ))
            collapsed.append(active_tasks[0])
            continue

        collapsed.append(max(
            project_tasks,
            key=lambda task: task.completed_at or task.updated_at,
        ))

    return sorted(
        collapsed,
        key=lambda task: (
            task.status == FollowUpTask.Status.COMPLETED,
            task.due_at or now + timedelta(days=36500),
            -task.updated_at.timestamp(),
        ),
    )


@login_required
@require_http_methods(['GET', 'POST'])
def follow_up_tasks(request):
    if request.method == 'POST':
        payload, error_response = parse_payload(request)
        if error_response:
            return error_response

        project_id = payload.get('projectId')
        if not str(project_id or '').isdigit():
            return JsonResponse({'error': '请选择关联项目'}, status=400)
        project = get_object_or_404(
            Project.objects.select_related(
                'customer',
                'business_owner__employee_profile',
                'technical_owner__employee_profile',
            ),
            id=int(project_id),
            is_active=True,
        )
        if project.customer.grade not in FOLLOW_UP_CUSTOMER_GRADES:
            return JsonResponse({'error': '只有 A、B 类客户可以新增跟进任务'}, status=400)
        if not can_manage_project(request.user, project):
            return JsonResponse({'error': '无权为该项目新增任务'}, status=403)

        title = str(payload.get('title', '')).strip()
        if not title:
            return JsonResponse({'error': '请填写任务内容'}, status=400)
        if len(title) > 200:
            return JsonResponse({'error': '任务内容不能超过 200 个字'}, status=400)

        role = str(payload.get('role', '')).strip()
        if role not in FollowUpTask.Role.values:
            return JsonResponse({'error': '请选择负责岗位'}, status=400)
        raw_due_at = str(payload.get('dueAt', '')).strip()
        due_at = parse_task_due_at(raw_due_at)
        if raw_due_at and due_at is None:
            return JsonResponse({'error': '计划完成时间格式不正确'}, status=400)

        assignee = (
            project.business_owner
            if role == FollowUpTask.Role.BUSINESS
            else project.technical_owner
        )
        task = FollowUpTask.objects.create(
            project=project,
            title=title,
            target_progress=project.progress,
            is_manual=True,
            role=role,
            assignee=assignee,
            due_at=due_at,
        )
        refreshed = FollowUpTask.objects.select_related(
            'assignee__employee_profile',
            'project__customer',
            'project__business_owner__employee_profile',
            'project__technical_owner__employee_profile',
        ).get(id=task.id)
        return JsonResponse({'task': serialize_follow_up_task(refreshed, request.user)}, status=201)

    tasks = FollowUpTask.objects.filter(
        project__customer__grade__in=FOLLOW_UP_CUSTOMER_GRADES,
    ).select_related(
        'assignee__employee_profile',
        'project__customer',
        'project__business_owner__employee_profile',
        'project__technical_owner__employee_profile',
    )
    can_view_all = request.user.is_superuser
    if can_view_all:
        employee_id = request.GET.get('employee')
        if employee_id and employee_id.isdigit():
            tasks = tasks.filter(assignee_id=int(employee_id))
    else:
        tasks = tasks.filter(assignee=request.user)

    now = timezone.now()
    task_list = collapse_tasks_by_project([
        task for task in tasks
        if follow_up_task_matches_project_progress(task)
    ], now)

    status = request.GET.get('status')
    if status == 'overdue':
        task_list = [
            task for task in task_list
            if task.status != FollowUpTask.Status.COMPLETED
            and task.due_at
            and task.due_at < now
        ]
    elif status in FollowUpTask.Status.values:
        task_list = [task for task in task_list if task.status == status]

    today_end = now.replace(hour=23, minute=59, second=59, microsecond=999999)
    summary = {
        'total': len(task_list),
        'pending': sum(
            1 for task in task_list
            if task.status == FollowUpTask.Status.PENDING
            and not (task.due_at and task.due_at < now)
        ),
        'today': sum(
            1 for task in task_list
            if task.status == FollowUpTask.Status.PENDING
            and task.due_at
            and now <= task.due_at <= today_end
        ),
        'overdue': sum(
            1 for task in task_list
            if task.status != FollowUpTask.Status.COMPLETED
            and task.due_at
            and task.due_at < now
        ),
        'waiting': sum(
            1 for task in task_list
            if task.status == FollowUpTask.Status.WAITING
            and not (task.due_at and task.due_at < now)
        ),
        'completed': sum(1 for task in task_list if task.status == FollowUpTask.Status.COMPLETED),
    }

    employees = []
    if can_view_all:
        users = (
            get_user_model()
            .objects.filter(is_active=True)
            .filter(
                Q(employee_profile__role__in=(
                    EmployeeProfile.Role.BUSINESS,
                    EmployeeProfile.Role.TECHNICAL,
                    EmployeeProfile.Role.BOTH,
                ))
                | Q(is_superuser=True)
            )
            .select_related('employee_profile')
            .distinct()
            .order_by('username')
        )
        employees = [
            {
                'id': user.id,
                'name': employee_display_name(user),
                'username': user.username,
            }
            for user in users
        ]

    return JsonResponse({
        'tasks': [serialize_follow_up_task(task, request.user) for task in task_list],
        'summary': summary,
        'canViewAll': can_view_all,
        'employees': employees,
    })


@login_required
@require_http_methods(['PATCH', 'DELETE'])
def follow_up_task_detail(request, task_id):
    task = get_object_or_404(
        FollowUpTask.objects.select_related(
            'project__customer',
            'project__business_owner__employee_profile',
            'project__technical_owner__employee_profile',
        ),
        id=task_id,
    )
    if task.project.customer.grade not in FOLLOW_UP_CUSTOMER_GRADES:
        return JsonResponse({'error': 'C、D 类客户不启用跟进记录功能'}, status=403)
    if not can_manage_follow_up(request.user, task):
        return JsonResponse({'error': '无权修改或删除该任务'}, status=403)

    if request.method == 'DELETE':
        deleted_id = task.id
        task.delete()
        return JsonResponse({'deleted': True, 'taskId': deleted_id})

    payload, error_response = parse_payload(request)
    if error_response:
        return error_response

    title = str(payload.get('title', task.title)).strip()
    if not title:
        return JsonResponse({'error': '请填写任务内容'}, status=400)
    if len(title) > 200:
        return JsonResponse({'error': '任务内容不能超过 200 个字'}, status=400)

    role = str(payload.get('role', task.role)).strip()
    if role not in FollowUpTask.Role.values:
        return JsonResponse({'error': '请选择负责岗位'}, status=400)

    due_at = task.due_at
    if 'dueAt' in payload:
        raw_due_at = str(payload.get('dueAt', '')).strip()
        due_at = parse_task_due_at(raw_due_at)
        if raw_due_at and due_at is None:
            return JsonResponse({'error': '计划完成时间格式不正确'}, status=400)

    previous_status = task.status
    status = str(payload.get('status', task.status)).strip()
    if status not in FollowUpTask.Status.values:
        return JsonResponse({'error': '任务状态不正确'}, status=400)
    result = str(payload.get('result', task.result)).strip()
    if status == FollowUpTask.Status.COMPLETED and not result:
        return JsonResponse({'error': '完成任务前请填写跟进结果'}, status=400)

    status_changed = status != previous_status
    transitioned_to_completed = (
        status_changed and status == FollowUpTask.Status.COMPLETED
    )
    task.title = title
    task.role = role
    task.assignee = (
        task.project.business_owner
        if role == FollowUpTask.Role.BUSINESS
        else task.project.technical_owner
    )
    task.due_at = due_at
    task.status = status
    task.result = result
    if status_changed:
        task.completed_at = (
            timezone.now()
            if status == FollowUpTask.Status.COMPLETED
            else None
        )
    task.save()

    if transitioned_to_completed and not task.is_manual:
        project = task.project
        project.progress = task.target_progress
        project.save()
        customer = project.customer
        primary_project_id = customer.projects.order_by('created_at', 'id').values_list('id', flat=True).first()
        if primary_project_id == project.id:
            customer.progress = project.progress
            customer.save(update_fields=('progress', 'updated_at'))

    refreshed = FollowUpTask.objects.select_related(
        'assignee__employee_profile',
        'project__customer',
        'project__business_owner__employee_profile',
        'project__technical_owner__employee_profile',
    ).get(id=task.id)
    return JsonResponse({'task': serialize_follow_up_task(refreshed, request.user)})


def serialize_tomorrow_item(item):
    return {
        'id': item.id,
        'title': item.title,
        'note': item.note,
        'plannedDate': item.planned_date.isoformat(),
        'isCompleted': item.is_completed,
        'completedAt': item.completed_at.isoformat() if item.completed_at else '',
        'createdAt': item.created_at.isoformat(),
    }


def parse_planned_date(value):
    if not value:
        return timezone.localdate() + timedelta(days=1)
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return None


@login_required
@require_http_methods(['GET', 'POST'])
def tomorrow_items(request):
    if request.method == 'GET':
        planned_date = parse_planned_date(request.GET.get('date'))
        if planned_date is None:
            return JsonResponse({'error': '计划日期格式不正确'}, status=400)
        items = TomorrowItem.objects.filter(
            user=request.user,
            planned_date=planned_date,
        )
        return JsonResponse({
            'date': planned_date.isoformat(),
            'items': [serialize_tomorrow_item(item) for item in items],
        })

    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    title = str(payload.get('title', '')).strip()
    note = str(payload.get('note', '')).strip()
    planned_date = parse_planned_date(payload.get('plannedDate'))
    if not title:
        return JsonResponse({'error': '请输入事项内容'}, status=400)
    if len(title) > 160:
        return JsonResponse({'error': '事项内容不能超过160个字'}, status=400)
    if planned_date is None:
        return JsonResponse({'error': '计划日期格式不正确'}, status=400)
    item = TomorrowItem.objects.create(
        user=request.user,
        title=title,
        note=note,
        planned_date=planned_date,
    )
    return JsonResponse({'item': serialize_tomorrow_item(item)}, status=201)


@login_required
@require_http_methods(['PATCH', 'DELETE'])
def tomorrow_item_detail(request, item_id):
    item = get_object_or_404(TomorrowItem, id=item_id, user=request.user)
    if request.method == 'DELETE':
        item.delete()
        return JsonResponse({'deleted': True})

    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    update_fields = []
    if 'title' in payload:
        title = str(payload.get('title', '')).strip()
        if not title:
            return JsonResponse({'error': '请输入事项内容'}, status=400)
        if len(title) > 160:
            return JsonResponse({'error': '事项内容不能超过160个字'}, status=400)
        item.title = title
        update_fields.append('title')
    if 'note' in payload:
        item.note = str(payload.get('note', '')).strip()
        update_fields.append('note')
    if 'isCompleted' in payload:
        item.is_completed = bool(payload.get('isCompleted'))
        item.completed_at = timezone.now() if item.is_completed else None
        update_fields.extend(('is_completed', 'completed_at'))
    if update_fields:
        item.save(update_fields=(*dict.fromkeys(update_fields), 'updated_at'))
    return JsonResponse({'item': serialize_tomorrow_item(item)})


def serialize_project_type(project_type):
    active_projects = project_type.projects.filter(
        is_active=True,
        customer__grade__in=('A', 'B', 'C'),
    ).count()
    return {
        'id': project_type.id,
        'name': project_type.name,
        'isActive': project_type.is_active,
        'projectCount': active_projects,
    }


@login_required
@require_http_methods(['GET', 'POST'])
def project_types(request):
    if request.method == 'GET':
        items = ProjectType.objects.all()
        return JsonResponse({
            'types': [serialize_project_type(item) for item in items],
            'canManage': request.user.is_superuser,
        })

    if not request.user.is_superuser:
        return JsonResponse({'error': '只有超级管理员可以新增项目类型'}, status=403)
    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    name = str(payload.get('name', '')).strip()
    if not name:
        return JsonResponse({'error': '请输入项目类型名称'}, status=400)
    if len(name) > 50:
        return JsonResponse({'error': '项目类型名称不能超过50个字'}, status=400)
    item = ProjectType.objects.filter(name=name).first()
    if item:
        if item.is_active:
            return JsonResponse({'error': '该项目类型已经存在'}, status=400)
        item.is_active = True
        item.save(update_fields=('is_active', 'updated_at'))
    else:
        item = ProjectType.objects.create(name=name)
    return JsonResponse({'projectType': serialize_project_type(item)}, status=201)


@login_required
@require_http_methods(['DELETE'])
def project_type_detail(request, project_type_id):
    if not request.user.is_superuser:
        return JsonResponse({'error': '只有超级管理员可以停用项目类型'}, status=403)
    item = get_object_or_404(ProjectType, id=project_type_id)
    item.is_active = False
    item.save(update_fields=('is_active', 'updated_at'))
    return JsonResponse({'projectType': serialize_project_type(item)})


def visible_progress_updates(user):
    queryset = ProjectProgressUpdate.objects.filter(
        project__is_active=True,
        project__customer__grade__in=FOLLOW_UP_CUSTOMER_GRADES,
    )
    if not user.is_superuser:
        queryset = queryset.filter(
            Q(project__business_owner=user) | Q(project__technical_owner=user)
        )
    return queryset.distinct()


def default_progress_interval(project):
    current_progress = normalize_progress(project.progress)
    try:
        current_index = PROGRESS_STAGES.index(current_progress)
    except ValueError:
        current_index = 0
    from_index = min(current_index, len(PROGRESS_STAGES) - 2)
    return PROGRESS_STAGES[from_index], PROGRESS_STAGES[from_index + 1]


def progress_interval_from_payload(payload, project, item=None):
    default_from, default_to = default_progress_interval(project)
    from_progress = str(
        payload.get('fromProgress', item.from_progress if item else default_from)
    ).strip()
    to_progress = str(
        payload.get('toProgress', item.to_progress if item else default_to)
    ).strip()
    try:
        from_index = PROGRESS_STAGES.index(from_progress)
        to_index = PROGRESS_STAGES.index(to_progress)
    except ValueError:
        return None, '所属阶段区间不正确'
    if to_index != from_index + 1:
        return None, '进展记录必须属于两个相邻的正式阶段之间'

    raw_occurred_at = str(payload.get('occurredAt', '')).strip()
    if raw_occurred_at:
        occurred_at = parse_datetime(raw_occurred_at)
        if occurred_at is None:
            return None, '记录时间格式不正确'
        if timezone.is_naive(occurred_at):
            occurred_at = timezone.make_aware(occurred_at, timezone.get_current_timezone())
    else:
        occurred_at = item.occurred_at if item else timezone.now()
    return {
        'from_progress': from_progress,
        'to_progress': to_progress,
        'occurred_at': occurred_at,
    }, ''


def progress_update_nature(item):
    """Infer the responsibility represented by both new and historical updates."""
    creator = item.created_by
    profile = getattr(creator, 'employee_profile', None) if creator else None
    role = (
        EmployeeProfile.Role.BOTH if creator and creator.is_superuser
        else (profile.role if profile else EmployeeProfile.Role.OTHER)
    )
    if role == EmployeeProfile.Role.BUSINESS:
        return EmployeeProfile.Role.BUSINESS
    if role == EmployeeProfile.Role.TECHNICAL:
        return EmployeeProfile.Role.TECHNICAL
    if role == EmployeeProfile.Role.BOTH:
        is_business_owner = item.project.business_owner_id == item.created_by_id
        is_technical_owner = item.project.technical_owner_id == item.created_by_id
        if is_business_owner and not is_technical_owner:
            return EmployeeProfile.Role.BUSINESS
        if is_technical_owner and not is_business_owner:
            return EmployeeProfile.Role.TECHNICAL
        return EmployeeProfile.Role.BOTH
    if item.project.technical_owner_id == item.created_by_id:
        return EmployeeProfile.Role.TECHNICAL
    return EmployeeProfile.Role.BUSINESS


def progress_update_group(item, nature):
    if nature == EmployeeProfile.Role.TECHNICAL:
        owner = item.project.technical_owner
        profile = getattr(owner, 'employee_profile', None) if owner else None
        return {
            'key': f'technical-{item.project.technical_owner_id or "unassigned"}',
            'nature': nature,
            'natureLabel': '技术',
            'ownerId': item.project.technical_owner_id or '',
            'ownerName': employee_display_name(owner),
            'level': profile.technical_level if profile else (5 if owner and owner.is_superuser else 1),
            'levelLabel': f'技术等级 {profile.technical_level if profile else (5 if owner and owner.is_superuser else 1)} 级',
        }
    if nature == EmployeeProfile.Role.BOTH:
        creator = item.created_by
        profile = getattr(creator, 'employee_profile', None) if creator else None
        business_level = profile.business_level if profile else (5 if creator and creator.is_superuser else 1)
        technical_level = profile.technical_level if profile else (5 if creator and creator.is_superuser else 1)
        return {
            'key': f'both-{item.created_by_id or "system"}',
            'nature': nature,
            'natureLabel': '商务与技术',
            'ownerId': item.created_by_id or '',
            'ownerName': employee_display_name(creator, '系统'),
            'level': max(business_level, technical_level),
            'levelLabel': f'商务 {business_level} 级 · 技术 {technical_level} 级',
        }
    owner = item.project.business_owner
    profile = getattr(owner, 'employee_profile', None) if owner else None
    return {
        'key': f'business-{item.project.business_owner_id or "unassigned"}',
        'nature': EmployeeProfile.Role.BUSINESS,
        'natureLabel': '商务',
        'ownerId': item.project.business_owner_id or '',
        'ownerName': employee_display_name(owner),
        'level': profile.business_level if profile else (5 if owner and owner.is_superuser else 1),
        'levelLabel': f'商务等级 {profile.business_level if profile else (5 if owner and owner.is_superuser else 1)} 级',
    }


def serialize_progress_update(item, user=None):
    created_time = timezone.localtime(item.occurred_at)
    actual_created_time = timezone.localtime(item.created_at)
    updated_time = timezone.localtime(item.updated_at)
    nature = progress_update_nature(item)
    update_group = progress_update_group(item, nature)
    return {
        'id': item.id,
        'content': item.content,
        'fromProgress': item.from_progress,
        'toProgress': item.to_progress,
        'intervalLabel': f'{item.from_progress} → {item.to_progress}',
        'occurredAt': item.occurred_at.isoformat(),
        'createdAt': item.created_at.isoformat(),
        'createdLabel': created_time.strftime('%Y年%m月%d日 %H:%M'),
        'createdDateLabel': created_time.strftime('%m月%d日'),
        'createdTimeLabel': created_time.strftime('%H:%M'),
        'actualCreatedLabel': actual_created_time.strftime('%Y年%m月%d日 %H:%M'),
        'updatedAt': item.updated_at.isoformat(),
        'updatedLabel': updated_time.strftime('%Y年%m月%d日 %H:%M'),
        'updatedDateLabel': updated_time.strftime('%m月%d日'),
        'updatedTimeLabel': updated_time.strftime('%H:%M'),
        'createdBy': employee_display_name(item.created_by, '系统'),
        'updateNature': nature,
        'updateNatureLabel': update_group['natureLabel'],
        'updateGroup': update_group,
        'canEdit': bool(user and (
            can_manage_project(user, item.project)
            or item.created_by_id == user.id
        )),
        'canDelete': bool(user and (
            can_manage_project(user, item.project)
            or item.created_by_id == user.id
        )),
        'project': serialize_project(item.project),
    }


@login_required
@require_http_methods(['GET', 'POST'])
def progress_updates(request):
    if request.method == 'POST':
        payload, error_response = parse_payload(request)
        if error_response:
            return error_response
        project = get_object_or_404(
            Project.objects.select_related(
                'customer', 'project_type',
                'business_owner__employee_profile',
                'technical_owner__employee_profile',
            ),
            id=payload.get('projectId'),
            is_active=True,
        )
        if project.customer.grade not in FOLLOW_UP_CUSTOMER_GRADES:
            return JsonResponse({'error': '只有 A、B 类客户可以新增进度更新'}, status=400)
        if not can_manage_project(request.user, project):
            return JsonResponse({'error': '只有该项目负责人或超级管理员可以更新进度'}, status=403)
        content = str(payload.get('content', '')).strip()
        if not content:
            return JsonResponse({'error': '请填写本次项目进展'}, status=400)
        if len(content) > 500:
            return JsonResponse({'error': '进度更新不能超过500个字'}, status=400)
        interval, interval_error = progress_interval_from_payload(payload, project)
        if interval_error:
            return JsonResponse({'error': interval_error}, status=400)
        item = ProjectProgressUpdate.objects.create(
            project=project,
            content=content,
            created_by=request.user,
            **interval,
        )
        return JsonResponse({'update': serialize_progress_update(item, request.user)}, status=201)

    start_at = timezone.now() - timedelta(days=7)
    base_queryset = visible_progress_updates(request.user).select_related(
        'created_by__employee_profile',
        'project__customer',
        'project__project_type',
        'project__business_owner__employee_profile',
        'project__technical_owner__employee_profile',
    )
    recent_items = list(base_queryset.filter(
        occurred_at__gte=start_at,
    ).order_by('-occurred_at', '-id'))
    read_ids = set(ProgressUpdateReadReceipt.objects.filter(
        user=request.user,
    ).values_list('progress_update_id', flat=True))
    items = [item for item in recent_items if item.id not in read_ids]
    read_items = list(base_queryset.filter(
        read_receipts__user=request.user,
    ).order_by('-occurred_at', '-id'))
    state = ProgressUpdateReadState.objects.filter(user=request.user).first()
    unread_count = sum(
        1 for item in recent_items
        if not state or not state.last_read_at or item.created_at > state.last_read_at
    )
    return JsonResponse({
        'updates': [serialize_progress_update(item, request.user) for item in items],
        'readUpdates': [serialize_progress_update(item, request.user) for item in read_items],
        'periodUpdates': [serialize_progress_update(item, request.user) for item in recent_items],
        'total': len(recent_items),
        'activeTotal': len(items),
        'unreadCount': unread_count,
        'periodStart': timezone.localtime(start_at).strftime('%Y-%m-%d'),
        'periodEnd': timezone.localdate().isoformat(),
    })


@login_required
@require_http_methods(['PATCH', 'DELETE'])
def progress_update_detail(request, update_id):
    item = get_object_or_404(
        ProjectProgressUpdate.objects.select_related(
            'created_by__employee_profile',
            'project__customer',
            'project__project_type',
            'project__business_owner__employee_profile',
            'project__technical_owner__employee_profile',
        ),
        id=update_id,
        project__is_active=True,
    )
    if not can_manage_project(request.user, item.project) and item.created_by_id != request.user.id:
        return JsonResponse({'error': '只有该项目负责人、记录人或超级管理员可以修改或删除进度'}, status=403)
    if request.method == 'DELETE':
        item_id = item.id
        item.delete()
        return JsonResponse({'deleted': True, 'updateId': item_id})
    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    content = str(payload.get('content', '')).strip()
    if not content:
        return JsonResponse({'error': '请填写本次项目进展'}, status=400)
    if len(content) > 500:
        return JsonResponse({'error': '进度更新不能超过500个字'}, status=400)
    interval, interval_error = progress_interval_from_payload(payload, item.project, item)
    if interval_error:
        return JsonResponse({'error': interval_error}, status=400)
    item.content = content
    item.from_progress = interval['from_progress']
    item.to_progress = interval['to_progress']
    item.occurred_at = interval['occurred_at']
    item.save(update_fields=('content', 'from_progress', 'to_progress', 'occurred_at', 'updated_at'))
    return JsonResponse({'update': serialize_progress_update(item, request.user)})


@login_required
@require_http_methods(['POST'])
def mark_progress_updates_read(request):
    state, _ = ProgressUpdateReadState.objects.get_or_create(user=request.user)
    state.last_read_at = timezone.now()
    state.save(update_fields=('last_read_at', 'updated_at'))
    return JsonResponse({'read': True, 'lastReadAt': state.last_read_at.isoformat()})


@login_required
@require_http_methods(['POST'])
def mark_progress_update_read(request, update_id):
    item = get_object_or_404(
        visible_progress_updates(request.user).select_related(
            'created_by__employee_profile',
            'project__customer',
            'project__project_type',
            'project__business_owner__employee_profile',
            'project__technical_owner__employee_profile',
        ),
        id=update_id,
    )
    receipt, _ = ProgressUpdateReadReceipt.objects.get_or_create(
        user=request.user,
        progress_update=item,
    )
    return JsonResponse({
        'read': True,
        'readAt': receipt.read_at.isoformat(),
        'update': serialize_progress_update(item, request.user),
    })


def combine_visitor_remarks(*values):
    parts = []
    for value in values:
        normalized = str(value or '').strip()
        if normalized and normalized not in parts:
            parts.append(normalized)
    return '\n'.join(parts)


def serialize_visitor_record(item, user=None):
    hosts = list(item.hosts.all())
    if not hosts and item.host:
        hosts = [item.host]
    visitor_contacts = item.visitor_contacts if isinstance(item.visitor_contacts, list) else []
    visitor_contacts = [
        {
            'name': str(contact.get('name', '')).strip(),
            'role': str(contact.get('role', '')).strip(),
        }
        for contact in visitor_contacts
        if isinstance(contact, dict) and str(contact.get('name', '')).strip()
    ]
    if not visitor_contacts and item.contact_name:
        visitor_contacts = [{'name': item.contact_name, 'role': ''}]
    return {
        'id': item.id,
        'visitDate': item.visit_date.isoformat(),
        'visitorCount': item.visitor_count,
        'contactName': item.contact_name,
        'visitorContacts': visitor_contacts,
        'purpose': item.purpose,
        # notes 是旧版“洽谈总结”字段。只读时并入备注，避免历史数据不可见。
        'remarks': combine_visitor_remarks(item.remarks, item.notes),
        'customer': {
            'id': item.customer_id,
            'name': item.customer.name,
            'grade': item.customer.grade,
            'phone': item.customer.phone,
            'source': item.customer.source,
        },
        'hostId': item.host_id or '',
        'hostName': employee_display_name(item.host),
        'hostIds': [host.id for host in hosts],
        'hostNames': [employee_display_name(host) for host in hosts],
        'createdBy': employee_display_name(item.created_by, '系统'),
        'updatedAt': item.updated_at.isoformat(),
        'canManage': can_manage_visitor_record(user, item) if user else False,
    }


def can_manage_visitor_record(user, item):
    if (
        user.is_superuser
        or item.created_by_id == user.id
        or item.host_id == user.id
        or any(host.id == user.id for host in item.hosts.all())
    ):
        return True
    profile = getattr(user, 'employee_profile', None)
    return bool(profile and profile.role in (EmployeeProfile.Role.BUSINESS, EmployeeProfile.Role.BOTH))


def normalize_visitor_contacts(payload):
    contacts = payload.get('visitorContacts')
    if contacts is None:
        legacy_name = str(payload.get('contactName', '')).strip()
        contacts = [{'name': legacy_name, 'role': ''}] if legacy_name else []
    if not isinstance(contacts, list):
        return '来访客户负责人格式不正确', []
    if len(contacts) > 50:
        return '单次最多登记 50 位来访客户负责人', []
    result = []
    for index, contact in enumerate(contacts, start=1):
        if not isinstance(contact, dict):
            return f'第 {index} 位来访客户负责人格式不正确', []
        name = str(contact.get('name', '')).strip()
        role = str(contact.get('role', '')).strip()
        if not name and not role:
            continue
        if not name:
            return f'请填写第 {index} 位来访客户负责人的姓名', []
        result.append({'name': name[:80], 'role': role[:100]})
    return '', result


def visitor_hosts_from_payload(payload):
    raw_host_ids = payload.get('hostIds')
    if raw_host_ids is None:
        raw_host_ids = [payload.get('hostId')] if payload.get('hostId') else []
    if not isinstance(raw_host_ids, list):
        return '接待人员格式不正确', []
    host_ids = []
    for raw_host_id in raw_host_ids:
        try:
            host_id = int(raw_host_id)
        except (TypeError, ValueError):
            return '所选接待人员账号不可用', []
        if host_id not in host_ids:
            host_ids.append(host_id)
    if len(host_ids) > 30:
        return '单次最多选择 30 位接待人员', []
    hosts = list(
        get_user_model().objects.filter(id__in=host_ids, is_active=True)
        .filter(Q(is_superuser=True) | Q(employee_profile__isnull=False))
        .select_related('employee_profile')
        .order_by('id')
        .distinct()
    )
    host_by_id = {host.id: host for host in hosts}
    if len(host_by_id) != len(host_ids):
        return '所选接待人员账号不可用', []
    return '', [host_by_id[host_id] for host_id in host_ids]


def update_visitor_from_payload(item, payload):
    customer = Customer.objects.filter(id=payload.get('customerId')).first()
    if not customer:
        return '请选择需要关联的客资', []
    visit_date = parse_date(str(payload.get('visitDate', '')))
    if not visit_date:
        return '请选择有效的来访日期', []
    try:
        visitor_count = int(payload.get('visitorCount', 1))
    except (TypeError, ValueError):
        return '来访人数必须是有效数字', []
    if visitor_count < 1 or visitor_count > 999:
        return '来访人数应在 1 至 999 人之间', []
    purpose = str(payload.get('purpose', '')).strip()
    if not purpose:
        return '请填写来访目的', []
    contacts_error, visitor_contacts = normalize_visitor_contacts(payload)
    if contacts_error:
        return contacts_error, []
    hosts_error, hosts = visitor_hosts_from_payload(payload)
    if hosts_error:
        return hosts_error, []
    item.customer = customer
    item.visit_date = visit_date
    item.visitor_count = max(visitor_count, len(visitor_contacts))
    item.visitor_contacts = visitor_contacts
    item.contact_name = '、'.join(contact['name'] for contact in visitor_contacts)[:100]
    item.purpose = purpose[:200]
    if 'remarks' in payload or 'notes' in payload:
        # 兼容旧客户端的 notes 入参；编辑后将两份内容收敛到唯一的备注字段。
        item.remarks = combine_visitor_remarks(payload.get('remarks'), payload.get('notes'))
        item.notes = ''
    item.host = hosts[0] if hosts else None
    return '', hosts


def visitor_record_queryset():
    return VisitorRecord.objects.select_related(
        'customer', 'host__employee_profile', 'created_by__employee_profile',
    ).prefetch_related('hosts__employee_profile')


def filtered_visitor_records(request, queryset=None):
    items = queryset if queryset is not None else visitor_record_queryset()
    keyword = request.GET.get('q', '').strip()
    if keyword:
        normalized_keyword = keyword.casefold()
        contact_record_ids = [
            record_id
            for record_id, contacts in items.values_list('id', 'visitor_contacts')
            if normalized_keyword in json.dumps(contacts, ensure_ascii=False).casefold()
        ]
        items = items.filter(
            Q(customer__name__icontains=keyword)
            | Q(contact_name__icontains=keyword)
            | Q(id__in=contact_record_ids)
            | Q(purpose__icontains=keyword)
            | Q(remarks__icontains=keyword)
            | Q(notes__icontains=keyword)
            | Q(hosts__username__icontains=keyword)
            | Q(hosts__first_name__icontains=keyword)
            | Q(hosts__last_name__icontains=keyword)
            | Q(hosts__employee_profile__nickname__icontains=keyword)
            | Q(host__username__icontains=keyword)
            | Q(host__employee_profile__nickname__icontains=keyword)
        )
    date_from = parse_date(request.GET.get('dateFrom', ''))
    date_to = parse_date(request.GET.get('dateTo', ''))
    if date_from:
        items = items.filter(visit_date__gte=date_from)
    if date_to:
        items = items.filter(visit_date__lte=date_to)
    source = request.GET.get('source', '').strip()
    if source:
        items = items.filter(customer__source=source)
    host_id = request.GET.get('hostId', '').strip()
    if host_id.isdigit():
        items = items.filter(Q(hosts__id=int(host_id)) | Q(host_id=int(host_id)))
    return items.distinct()


@login_required
@require_http_methods(['GET', 'POST'])
def visitor_records(request):
    if request.method == 'GET':
        items = filtered_visitor_records(request)
        return JsonResponse({'records': [serialize_visitor_record(item, request.user) for item in items]})

    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    item = VisitorRecord(created_by=request.user)
    validation_error, hosts = update_visitor_from_payload(item, payload)
    if validation_error:
        return JsonResponse({'error': validation_error}, status=400)
    item.save()
    item.hosts.set(hosts)
    item = visitor_record_queryset().get(id=item.id)
    return JsonResponse({'record': serialize_visitor_record(item, request.user)}, status=201)


VISITOR_IMPORT_HEADERS = (
    '序号',
    '考察时间',
    '关联的客资',
    '会谈内容',
    '备注',
    '客户来源（自媒体）',
    '接待人员',
)

VISITOR_EXPORT_HEADERS = (
    '序号',
    '考察时间',
    '关联的客资',
    '来访客户负责人',
    '会谈内容',
    '备注',
    '客户来源（自媒体）',
    '接待人员',
)


def parse_visitor_import_date(value):
    if isinstance(value, (int, float)) and 1 <= value <= 100000:
        return date(1899, 12, 30) + timedelta(days=int(value))
    raw_value = str(value or '').strip()
    parsed = parse_date(raw_value)
    if parsed:
        return parsed
    matched = re.fullmatch(r'(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?', raw_value)
    if matched:
        try:
            return date(*(int(part) for part in matched.groups()))
        except ValueError:
            return None
    return None


def visitor_import_host_map():
    result = {}
    users = get_user_model().objects.filter(is_active=True).select_related('employee_profile')
    for user in users:
        profile = getattr(user, 'employee_profile', None)
        if not user.is_superuser and not profile:
            continue
        names = {
            user.username.strip(),
            user.get_full_name().strip(),
            employee_display_name(user).strip(),
            (profile.nickname if profile else '').strip(),
        }
        for name in filter(None, names):
            result.setdefault(name, []).append(user)
    return result


def split_visitor_import_values(value):
    return [part.strip() for part in re.split(r'[、,，;；\n]+', str(value or '')) if part.strip()]


def parse_visitor_import_contacts(value):
    contacts = []
    for entry in [part.strip() for part in re.split(r'[;；\n]+', str(value or '')) if part.strip()]:
        parts = re.split(r'[|｜]', entry, maxsplit=1)
        contacts.append({
            'name': parts[0].strip()[:80],
            'role': parts[1].strip()[:100] if len(parts) > 1 else '',
        })
    return contacts


def resolve_visitor_import_hosts(value, host_map, row_number):
    result = []
    for host_name in split_visitor_import_values(value):
        if host_name == '待分配':
            continue
        hosts = host_map.get(host_name, [])
        if not hosts:
            return [], f'第 {row_number} 行未找到接待人员“{host_name}”'
        if len(hosts) > 1:
            return [], f'第 {row_number} 行接待人员“{host_name}”重名，请改填登录用户名'
        if hosts[0] not in result:
            result.append(hosts[0])
    return result, ''


@login_required
@require_http_methods(['POST'])
@transaction.atomic
def visitor_records_import(request):
    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    records = payload.get('records')
    if not isinstance(records, list) or not records:
        return JsonResponse({'error': 'Excel 中没有可导入的来访记录'}, status=400)
    if len(records) > 2000:
        return JsonResponse({'error': '单次最多导入 2000 条来访记录'}, status=400)
    first_record = records[0] if isinstance(records[0], dict) else {}
    missing_headers = [header for header in VISITOR_IMPORT_HEADERS if header not in first_record]
    if missing_headers:
        return JsonResponse({'error': f"缺少列：{'、'.join(missing_headers)}"}, status=400)

    customer_map = {}
    for customer in Customer.objects.all():
        customer_map.setdefault(customer.name.strip(), []).append(customer)
    host_map = visitor_import_host_map()
    pending_items = []
    for row_number, row in enumerate(records, start=2):
        if not isinstance(row, dict):
            transaction.set_rollback(True)
            return JsonResponse({'error': f'第 {row_number} 行格式不正确'}, status=400)
        customer_name = str(row.get('关联的客资') or '').strip()
        customers = customer_map.get(customer_name, [])
        if not customer_name:
            transaction.set_rollback(True)
            return JsonResponse({'error': f'第 {row_number} 行未填写关联的客资'}, status=400)
        if not customers:
            transaction.set_rollback(True)
            return JsonResponse({'error': f'第 {row_number} 行未找到客资“{customer_name}”'}, status=400)
        if len(customers) > 1:
            transaction.set_rollback(True)
            return JsonResponse({'error': f'第 {row_number} 行客资“{customer_name}”名称重复，无法确定关联对象'}, status=400)
        visit_date = parse_visitor_import_date(row.get('考察时间'))
        if not visit_date:
            transaction.set_rollback(True)
            return JsonResponse({'error': f'第 {row_number} 行考察时间格式不正确'}, status=400)
        purpose = str(row.get('会谈内容') or '').strip()
        if not purpose:
            transaction.set_rollback(True)
            return JsonResponse({'error': f'第 {row_number} 行未填写会谈内容'}, status=400)
        hosts, host_error = resolve_visitor_import_hosts(row.get('接待人员'), host_map, row_number)
        if host_error:
            transaction.set_rollback(True)
            return JsonResponse({'error': host_error}, status=400)
        visitor_contacts = parse_visitor_import_contacts(row.get('来访客户负责人'))
        pending_items.append((VisitorRecord(
            customer=customers[0],
            visit_date=visit_date,
            visitor_count=max(1, len(visitor_contacts)),
            visitor_contacts=visitor_contacts,
            contact_name='、'.join(contact['name'] for contact in visitor_contacts)[:100],
            purpose=purpose[:200],
            # 旧模板若仍带“洽谈总结”列，也无损并入备注。
            remarks=combine_visitor_remarks(row.get('备注'), row.get('洽谈总结')),
            notes='',
            host=hosts[0] if hosts else None,
            created_by=request.user,
        ), hosts))

    for item, hosts in pending_items:
        item.save()
        item.hosts.set(hosts)
    return JsonResponse({'imported': True, 'createdCount': len(pending_items)}, status=201)


@login_required
@require_http_methods(['GET'])
def visitor_records_export(request):
    items = list(filtered_visitor_records(request))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '客户来访接待'
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A2'
    widths = (9, 14, 24, 30, 34, 42, 20, 26)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    header_fill = PatternFill('solid', fgColor='0F7658')
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    body_font = Font(name='微软雅黑', size=10, color='193B33')
    border = Border(*(Side(style='thin', color='BDD5CC') for _ in range(4)))
    for column, header in enumerate(VISITOR_EXPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    sheet.row_dimensions[1].height = 32
    for index, item in enumerate(items, start=1):
        row = index + 1
        contacts = item.visitor_contacts if isinstance(item.visitor_contacts, list) else []
        contact_lines = [
            f"{contact.get('name', '').strip()}（{contact.get('role', '').strip()}）"
            if contact.get('role', '').strip() else contact.get('name', '').strip()
            for contact in contacts
            if isinstance(contact, dict) and contact.get('name', '').strip()
        ]
        if not contact_lines and item.contact_name:
            contact_lines = [item.contact_name]
        hosts = list(item.hosts.all())
        if not hosts and item.host:
            hosts = [item.host]
        values = (
            index,
            item.visit_date,
            item.customer.name,
            '\n'.join(contact_lines),
            item.purpose,
            combine_visitor_remarks(item.remarks, item.notes),
            item.customer.source,
            '\n'.join(employee_display_name(host) for host in hosts) or '待分配',
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = body_font
            cell.fill = PatternFill('solid', fgColor='FFFFFF' if index % 2 else 'F5FAF7')
            cell.alignment = Alignment(
                horizontal='center' if column in (1, 2, 7, 8) else 'left',
                vertical='center',
                wrap_text=True,
            )
            cell.border = border
        sheet.cell(row=row, column=2).number_format = 'yyyy-mm-dd'
        sheet.row_dimensions[row].height = max(34, 18 * max(len(contact_lines), len(hosts), 1))
    sheet.auto_filter.ref = f'A1:H{max(len(items) + 1, 1)}'
    output = BytesIO()
    workbook.save(output)
    filename = f'客户来访接待表_{timezone.localdate():%Y-%m-%d}.xlsx'
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = content_disposition_header(True, filename)
    response['Content-Length'] = output.tell()
    return response


@login_required
@require_http_methods(['PATCH', 'DELETE'])
def visitor_record_detail(request, record_id):
    item = get_object_or_404(
        visitor_record_queryset(),
        id=record_id,
    )
    if not can_manage_visitor_record(request.user, item):
        return JsonResponse({'error': '只有商务人员、接待负责人、记录人或超级管理员可以修改来访记录'}, status=403)
    if request.method == 'DELETE':
        item.delete()
        return JsonResponse({'deleted': True})
    payload, error_response = parse_payload(request)
    if error_response:
        return error_response
    validation_error, hosts = update_visitor_from_payload(item, payload)
    if validation_error:
        return JsonResponse({'error': validation_error}, status=400)
    item.save()
    item.hosts.set(hosts)
    item = visitor_record_queryset().get(id=item.id)
    return JsonResponse({'record': serialize_visitor_record(item, request.user)})
