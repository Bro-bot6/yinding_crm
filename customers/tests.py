import json
from datetime import date, datetime, time, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .admin import EmployeeUserCreationForm
from .models import (
    AgencyAuthorization,
    BusinessContract,
    Customer,
    CustomerProjectAssociation,
    EmployeeProfile,
    FollowUpTask,
    MaterialExperiment,
    PartnershipIdentity,
    ProgressUpdateReadReceipt,
    Project,
    ProjectProgressUpdate,
    ProjectType,
    SchemeCalculation,
    TomorrowItem,
    VisitorRecord,
)
from .services import PROGRESS_STAGES, NEXT_TASKS, normalize_progress


def seed_historical_stage_tasks(project):
    """Explicit legacy fixture: production no longer generates stage tasks."""
    current = PROGRESS_STAGES.index(normalize_progress(project.progress))
    if current == len(PROGRESS_STAGES) - 1:
        return
    target = PROGRESS_STAGES[current + 1]
    title, roles = NEXT_TASKS[target]
    for role in roles:
        FollowUpTask.objects.get_or_create(
            project=project, target_progress=target, role=role,
            defaults={'title': title, 'assignee': project.business_owner if role == 'business' else project.technical_owner},
        )


class ProjectTypeAndProgressUpdateApiTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username='overview-admin',
            password='123456',
        )
        self.employee = get_user_model().objects.create_user(
            username='overview-business',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.employee,
            role=EmployeeProfile.Role.BUSINESS,
        )
        self.project_type = ProjectType.objects.get(name='道路')
        self.customer = Customer.objects.create(
            name='测试 A 类客户',
            grade='A',
            progress='需求对接',
            business_owner=self.employee,
        )
        self.project = Project.objects.create(
            customer=self.customer,
            name='道路项目',
            project_type=self.project_type,
            business_owner=self.employee,
        )

    def test_only_superuser_can_manage_project_types(self):
        self.client.force_login(self.employee)
        denied = self.client.post(
            reverse('project-types'),
            data=json.dumps({'name': '厂房地坪'}),
            content_type='application/json',
        )
        self.assertEqual(denied.status_code, 403)

        self.client.force_login(self.admin)
        created = self.client.post(
            reverse('project-types'),
            data=json.dumps({'name': '厂房地坪'}),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 201)
        archived = self.client.delete(
            reverse('project-type-detail', args=(created.json()['projectType']['id'],)),
        )
        self.assertEqual(archived.status_code, 200)
        self.assertFalse(archived.json()['projectType']['isActive'])

    def test_ab_project_owner_can_add_update_and_unread_can_be_cleared(self):
        self.client.force_login(self.employee)
        created = self.client.post(
            reverse('progress-updates'),
            data=json.dumps({'projectId': self.project.id, 'content': '完成现场测量，明日整理参数。'}),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 201)
        self.assertEqual(ProjectProgressUpdate.objects.count(), 1)

        overview = self.client.get(reverse('progress-updates')).json()
        self.assertEqual(overview['total'], 1)
        self.assertEqual(overview['unreadCount'], 1)
        self.client.post(reverse('mark-progress-updates-read'))
        self.assertEqual(self.client.get(reverse('progress-updates')).json()['unreadCount'], 0)

    def test_progress_updates_are_grouped_by_business_or_technical_nature(self):
        technical = get_user_model().objects.create_user(
            username='overview-technical',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=technical,
            role=EmployeeProfile.Role.TECHNICAL,
            technical_level=4,
        )
        self.project.technical_owner = technical
        self.project.save(update_fields=('technical_owner',))
        ProjectProgressUpdate.objects.create(
            project=self.project,
            content='商务沟通更新',
            created_by=self.employee,
        )
        technical_update = ProjectProgressUpdate.objects.create(
            project=self.project,
            content='技术方案更新',
            created_by=technical,
        )

        self.client.force_login(self.admin)
        updates = self.client.get(reverse('progress-updates')).json()['updates']
        payload = {item['id']: item for item in updates}

        self.assertEqual(payload[technical_update.id]['updateNature'], 'technical')
        self.assertEqual(payload[technical_update.id]['updateGroup']['natureLabel'], '技术')
        self.assertEqual(payload[technical_update.id]['updateGroup']['level'], 4)
        business_update = next(item for item in updates if item['content'] == '商务沟通更新')
        self.assertEqual(business_update['updateGroup']['nature'], 'business')

    def test_project_update_can_be_read_from_customer_and_edited(self):
        item = ProjectProgressUpdate.objects.create(
            project=self.project,
            content='首次进度记录',
            created_by=self.employee,
        )
        self.client.force_login(self.employee)

        detail = self.client.get(reverse('customer-detail', args=(self.customer.id,)))
        self.assertEqual(detail.status_code, 200)
        record = detail.json()['progressUpdates'][0]
        self.assertEqual(record['content'], '首次进度记录')
        self.assertTrue(record['canEdit'])
        self.assertIn('createdLabel', record)
        self.assertIn('updatedLabel', record)

        edited = self.client.patch(
            reverse('progress-update-detail', args=(item.id,)),
            data=json.dumps({'content': '已补充现场测量数据'}),
            content_type='application/json',
        )
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(edited.json()['update']['content'], '已补充现场测量数据')
        item.refresh_from_db()
        self.assertEqual(item.content, '已补充现场测量数据')

        deleted = self.client.delete(reverse('progress-update-detail', args=(item.id,)))
        self.assertEqual(deleted.status_code, 200)
        self.assertTrue(deleted.json()['deleted'])
        self.assertFalse(ProjectProgressUpdate.objects.filter(id=item.id).exists())

    def test_small_progress_node_keeps_formal_stage_and_validates_adjacent_interval(self):
        self.client.force_login(self.employee)
        occurred_at = '2026-08-16T09:35:00+08:00'
        created = self.client.post(
            reverse('progress-updates'),
            data=json.dumps({
                'projectId': self.project.id,
                'content': '已协调寄样时间，等待物流单号。',
                'fromProgress': '技术验证',
                'toProgress': '客户深度沟通',
                'occurredAt': occurred_at,
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 201)
        payload = created.json()['update']
        self.assertEqual(payload['intervalLabel'], '技术验证 → 客户深度沟通')
        self.assertTrue(payload['occurredAt'].startswith('2026-08-16T09:35:00'))
        self.customer.refresh_from_db()
        self.project.refresh_from_db()
        self.assertEqual(self.customer.progress, '需求对接')
        self.assertEqual(self.project.progress, '需求对接')

        invalid = self.client.post(
            reverse('progress-updates'),
            data=json.dumps({
                'projectId': self.project.id,
                'content': '错误的跨阶段区间',
                'fromProgress': '需求对接',
                'toProgress': '方案与报价',
            }),
            content_type='application/json',
        )
        self.assertEqual(invalid.status_code, 400)

    def test_material_experiment_is_saved_and_returned_in_customer_detail(self):
        self.client.force_login(self.employee)
        saved = self.client.put(
            reverse('material-experiment', args=(self.customer.id,)),
            data=json.dumps({
                'stableMaterial': '现场素土',
                'day7Data': '无侧限抗压强度 2.1MPa',
                'day14Data': '无侧限抗压强度 3.4MPa',
                'day28Data': '待检测',
                'technicalMessage': '建议保持当前掺比继续观察。',
            }),
            content_type='application/json',
        )
        self.assertEqual(saved.status_code, 200)
        self.assertTrue(MaterialExperiment.objects.filter(project=self.project).exists())
        detail = self.client.get(reverse('customer-detail', args=(self.customer.id,)))
        self.assertEqual(detail.status_code, 200)
        experiment = detail.json()['materialExperiment']
        self.assertEqual(experiment['stableMaterial'], '现场素土')
        self.assertEqual(experiment['day28Data'], '待检测')
        self.assertTrue(experiment['canEdit'])
        self.assertIn('visitorRecords', detail.json())
        self.assertIn('schemeCalculation', detail.json())

    def test_update_can_be_archived_as_read_for_current_account(self):
        item = ProjectProgressUpdate.objects.create(
            project=self.project,
            content='等待领导查看的项目进展',
            created_by=self.employee,
        )
        self.client.force_login(self.employee)

        marked = self.client.post(reverse('mark-progress-update-read', args=(item.id,)))
        self.assertEqual(marked.status_code, 200)
        self.assertTrue(marked.json()['read'])
        self.assertTrue(ProgressUpdateReadReceipt.objects.filter(
            user=self.employee,
            progress_update=item,
        ).exists())

        overview = self.client.get(reverse('progress-updates')).json()
        self.assertEqual(overview['total'], 1)
        self.assertEqual(overview['updates'], [])
        self.assertEqual(overview['readUpdates'][0]['id'], item.id)
        self.assertEqual(overview['periodUpdates'][0]['id'], item.id)

    def test_c_customer_project_can_keep_small_progress_without_entering_ab_summary(self):
        self.customer.grade = 'C'
        self.customer.save(update_fields=('grade',))
        self.client.force_login(self.employee)
        response = self.client.post(
            reverse('progress-updates'),
            data=json.dumps({'projectId': self.project.id, 'content': 'C类项目日常跟进'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(ProjectProgressUpdate.objects.count(), 1)
        detail = self.client.get(
            reverse('customer-detail', args=(self.customer.id,)),
            {'projectId': self.project.id},
        ).json()
        self.assertEqual(detail['progressUpdates'][0]['content'], 'C类项目日常跟进')
        self.assertEqual(self.client.get(reverse('progress-updates')).json()['total'], 0)

    def test_customer_payload_updates_primary_project_type(self):
        self.client.force_login(self.admin)
        water = ProjectType.objects.get(name='水利')
        response = self.client.patch(
            reverse('customer-detail', args=(self.customer.id,)),
            data=json.dumps({
                'name': self.customer.name,
                'phone': '13800000021',
                'source': '抖音',
                'province': '浙江省',
                'city': '杭州市',
                'district': '西湖区',
                'grade': 'A',
                'description': '',
                'progress': '需求对接',
                'plan': '水利项目方案',
                'projectTypeId': water.id,
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.project.refresh_from_db()
        self.assertEqual(self.project.project_type, water)
        self.assertEqual(response.json()['customer']['projectTypeName'], '水利')


class OverviewStatisticsApiTests(TestCase):
    def setUp(self):
        Customer.objects.all().delete()
        self.admin = get_user_model().objects.create_superuser(
            username='overview-statistics-admin',
            password='123456',
        )
        self.business = get_user_model().objects.create_user(
            username='overview-statistics-business',
            password='123456',
        )
        self.other_business = get_user_model().objects.create_user(
            username='overview-statistics-other',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.business,
            role=EmployeeProfile.Role.BUSINESS,
        )
        EmployeeProfile.objects.create(
            user=self.other_business,
            role=EmployeeProfile.Role.BUSINESS,
        )
        self.project_type = ProjectType.objects.get(name='道路')

    def aware_at(self, day, hour=12):
        return timezone.make_aware(
            datetime.combine(day, time(hour=hour)),
            timezone.get_current_timezone(),
        )

    def test_overview_matches_persisted_records_and_exact_date_window(self):
        today = timezone.localdate()
        customers = [
            Customer.objects.create(
                name='今日 A 客资', grade='A', source='抖音', province='浙江省',
                business_owner=self.business,
            ),
            Customer.objects.create(
                name='四日前 B 客资', grade='B', source='视频号', province='',
                business_owner=self.business,
            ),
            Customer.objects.create(
                name='今日 C 客资', grade='C', source='朋友介绍', province='浙江省',
                business_owner=self.business,
            ),
            Customer.objects.create(
                name='今日第二条 A 客资', grade='A', source='抖音', province='江苏省',
                business_owner=self.business,
            ),
        ]
        Customer.objects.filter(id=customers[1].id).update(
            created_at=self.aware_at(today - timedelta(days=4)),
        )
        Customer.objects.create(
            name='不进入有效统计的 D 客资', grade='D', source='抖音', province='江苏省',
        )
        projects = [
            Project.objects.create(
                customer=customers[0], name='实施项目', project_type=self.project_type,
                progress='项目实施跟进', business_owner=self.business,
            ),
            Project.objects.create(
                customer=customers[1], name='已完成项目', project_type=None,
                progress='售后维护与需求挖掘', is_active=False,
                business_owner=self.business,
            ),
            Project.objects.create(
                customer=customers[2], name='C 类保留项目', project_type=self.project_type,
                progress='需求对接', business_owner=self.business,
            ),
            Project.objects.create(
                customer=customers[3], name='逾期任务项目', project_type=self.project_type,
                progress='需求对接', business_owner=self.business,
            ),
        ]
        VisitorRecord.objects.create(
            customer=customers[0], visit_date=today, visitor_count=3,
            purpose='现场考察', created_by=self.business,
        )
        VisitorRecord.objects.create(
            customer=customers[1], visit_date=today - timedelta(days=4), visitor_count=2,
            purpose='方案沟通', created_by=self.business,
        )
        due_dates = (
            self.aware_at(today, 23),
            self.aware_at(today + timedelta(days=1), 10),
            self.aware_at(today - timedelta(days=1), 10),
        )
        task_projects = (projects[0], projects[1], projects[3])
        for index, (project, due_at) in enumerate(zip(task_projects, due_dates)):
            FollowUpTask.objects.create(
                project=project,
                title=f'真实任务 {index}',
                target_progress=project.progress,
                is_manual=True,
                role=FollowUpTask.Role.BUSINESS,
                assignee=self.business,
                due_at=due_at,
            )
        ProjectProgressUpdate.objects.create(
            project=projects[0], content='近七天真实项目进展', created_by=self.business,
        )

        self.client.force_login(self.admin)
        response = self.client.get(reverse('overview-statistics'))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['period']['timezone'], 'Asia/Shanghai')
        self.assertEqual(payload['period']['startDate'], (today - timedelta(days=29)).isoformat())
        self.assertEqual(payload['period']['endDate'], today.isoformat())
        self.assertEqual(payload['customers']['total'], 4)
        self.assertEqual(payload['customers']['gradeCounts'], {'A': 2, 'B': 1, 'C': 1})
        self.assertEqual(sum(item['count'] for item in payload['customers']['sourceCounts']), 4)
        self.assertEqual(sum(item['total'] for item in payload['customers']['trendBuckets']), 4)
        self.assertEqual(payload['customers']['regions'], [
            {'name': '浙江省', 'count': 2},
            {'name': '江苏省', 'count': 1},
        ])
        self.assertEqual(payload['customers']['missingRegionCount'], 1)
        self.assertEqual(payload['projects']['total'], 4)
        self.assertEqual(payload['projects']['active'], 3)
        self.assertEqual(payload['projects']['completed'], 1)
        self.assertEqual(payload['projects']['implementation'], 1)
        self.assertEqual(payload['visitors']['visits'], 2)
        self.assertEqual(payload['visitors']['people'], 5)
        self.assertEqual(payload['followUps']['total'], 3)
        self.assertEqual(payload['followUps']['todayDue'], 1)
        self.assertEqual(payload['followUps']['tomorrowDue'], 1)
        self.assertEqual(payload['followUps']['overdue'], 1)
        self.assertEqual(payload['progressUpdates']['last7Days'], 1)

        self.client.force_login(self.business)
        employee_payload = self.client.get(reverse('overview-statistics')).json()
        self.assertEqual(employee_payload['followUps']['total'], 3)

    def test_overview_empty_state_returns_zero_without_fabricated_series(self):
        self.client.force_login(self.admin)

        payload = self.client.get(reverse('overview-statistics')).json()

        self.assertEqual(payload['customers']['total'], 0)
        self.assertEqual(payload['customers']['sourceCounts'], [
            {'name': '抖音', 'count': 0},
            {'name': '视频号', 'count': 0},
            {'name': '服务号', 'count': 0},
            {'name': '朋友介绍', 'count': 0},
        ])
        self.assertEqual(len(payload['customers']['trendBuckets']), 6)
        self.assertTrue(all(bucket['total'] == 0 for bucket in payload['customers']['trendBuckets']))
        self.assertEqual(payload['projects']['total'], 0)
        self.assertEqual(payload['projects']['typeCounts'], [])
        self.assertEqual(payload['visitors']['visits'], 0)
        self.assertTrue(all(bucket['visits'] == 0 and bucket['people'] == 0 for bucket in payload['visitors']['trendBuckets']))
        self.assertEqual(payload['followUps']['total'], 0)
        self.assertEqual(payload['progressUpdates']['last7Days'], 0)

    def test_overview_ignores_legacy_source_without_crashing(self):
        Customer.objects.create(
            name='历史其他来源客资',
            phone='legacy-source',
            grade='C',
            source='其他',
        )
        self.client.force_login(self.admin)

        response = self.client.get(reverse('overview-statistics'))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['customers']['total'], 1)
        self.assertEqual(payload['customers']['sourceCounts'], [
            {'name': '抖音', 'count': 0},
            {'name': '视频号', 'count': 0},
            {'name': '服务号', 'count': 0},
            {'name': '朋友介绍', 'count': 0},
        ])
        self.assertTrue(all(bucket['total'] == 0 for bucket in payload['customers']['trendBuckets']))

class SchemeCalculationApiTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username='scheme-admin',
            password='123456',
        )
        self.customer = Customer.objects.create(
            name='三层道路客户',
            grade='A',
            province='河北省',
            city='石家庄市',
        )
        self.project = Project.objects.create(
            customer=self.customer,
            name='园区道路项目',
            business_owner=self.admin,
        )
        self.payload = {
            'title': '园区道路三层方案造价分析',
            'layerCount': 3,
            'remarks': '需要先做试验段，运输费用另计。',
            'layers': [
                {
                    'name': '土凝岩稳定土', 'structureLayer': '面层',
                    'mixDescription': '土凝岩稳定土（PS-I）',
                    'length': 100, 'width': 4, 'thickness': 0.2,
                    'dosagePercent': 10, 'unitPrice': 600, 'density': 1.7,
                },
                {
                    'name': '土凝岩稳定土', 'structureLayer': '基层',
                    'mixDescription': '8%土凝岩稳定土（PS-I）',
                    'length': 100, 'width': 4, 'thickness': 0.2,
                    'dosagePercent': 8, 'unitPrice': 600, 'density': 1.7,
                },
                {
                    'name': '土凝岩稳定土', 'structureLayer': '底基层',
                    'mixDescription': '6%土凝岩稳定土（PS-I）',
                    'length': 100, 'width': 4, 'thickness': 0.2,
                    'dosagePercent': 6, 'unitPrice': 600, 'density': 1.7,
                },
            ],
        }

    def test_three_layer_calculation_is_saved_and_quantity_rounds_up(self):
        self.client.force_login(self.admin)
        response = self.client.put(
            reverse('scheme-calculation', args=(self.customer.id,)),
            data=json.dumps(self.payload),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        result = response.json()['calculation']
        self.assertEqual(result['layerCount'], 3)
        self.assertEqual(len(result['layers']), 3)
        self.assertEqual(result['layers'][0]['quantity'], 13.6)
        self.assertEqual(result['layers'][0]['totalPrice'], 8160)
        self.assertEqual(result['totalExactQuantity'], 32.64)
        self.assertEqual(result['totalQuantity'], 33)
        self.assertEqual(result['totalPrice'], 19584)
        self.assertEqual(result['layers'][0]['mixDescription'], '10%土凝岩稳定土（PS-I）')
        self.assertEqual(result['remarks'], self.payload['remarks'])
        self.assertTrue(SchemeCalculation.objects.filter(project=self.project).exists())

    def test_square_meter_price_uses_each_layers_actual_thickness(self):
        self.client.force_login(self.admin)
        self.payload['layers'][0]['thickness'] = 0.35
        response = self.client.put(
            reverse('scheme-calculation', args=(self.customer.id,)),
            data=json.dumps(self.payload),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        first_layer = response.json()['calculation']['layers'][0]
        self.assertEqual(first_layer['squareMeterPrice'], 35.7)

    def test_export_returns_downloadable_xlsx(self):
        self.client.force_login(self.admin)
        self.client.put(
            reverse('scheme-calculation', args=(self.customer.id,)),
            data=json.dumps(self.payload),
            content_type='application/json',
        )
        response = self.client.post(
            reverse('export-scheme-calculation', args=(self.customer.id,)),
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment;', response['Content-Disposition'])
        self.assertIn('.xlsx', response['Content-Disposition'])
        self.assertEqual(int(response['Content-Length']), len(response.content))
        self.assertGreater(len(response.content), 0)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet['J12'].value, '土凝岩单平米报价（元/㎡）')
        self.assertEqual(sheet.max_column, 12)
        self.assertEqual(sheet.cell(row=18, column=1).value, '报价备注 / TERMS & NOTES')
        self.assertEqual(sheet.cell(row=19, column=1).value, 1)
        self.assertEqual(sheet.cell(row=19, column=2).value, self.payload['remarks'])
        self.assertEqual(sheet.cell(row=20, column=1).value, 2)
        self.assertEqual(sheet.cell(row=21, column=1).value, 3)
        self.assertIsNone(sheet.auto_filter.ref)
        self.assertEqual(sheet['A8'].value, '推荐掺量')
        self.assertEqual(sheet['J8'].value, '材料总价')
        self.assertEqual(sheet['J9'].value, '¥19,584.00')
        self.assertEqual(sheet['G9'].value, '33 吨')
        self.assertEqual(sheet['K13'].value, 13.6)
        self.assertEqual(sheet['K16'].value, 33)
        self.assertEqual(sheet['A9'].value, '10% + 8% + 6%')
        self.assertEqual(sheet['D9'].value, '¥48.96 /㎡')
        self.assertEqual(sheet['D13'].value, '10%土凝岩稳定土（PS-I）')
        self.assertNotIn('说明：', ' '.join(
            str(cell.value or '') for row in sheet.iter_rows() for cell in row
        ))


class VisitorRecordApiTests(TestCase):
    def setUp(self):
        self.business = get_user_model().objects.create_user(
            username='visitor-business', password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.business,
            role=EmployeeProfile.Role.BUSINESS,
            business_level=4,
        )
        self.other = get_user_model().objects.create_user(
            username='visitor-other', password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.other,
            role=EmployeeProfile.Role.OTHER,
        )
        self.customer = Customer.objects.create(name='来访测试客户', grade='A')
        self.customer.source = '抖音'
        self.customer.save(update_fields=('source',))

    def visitor_payload(self):
        return {
            'customerId': self.customer.id,
            'visitDate': timezone.localdate().isoformat(),
            'visitorCount': 3,
            'visitorContacts': [
                {'name': '张伟', 'role': '项目经理'},
                {'name': '李敏', 'role': '项目总工'},
            ],
            'purpose': '参观样板工程',
            'remarks': '客户自驾到访',
            'hostIds': [self.business.id, self.other.id],
        }

    def test_create_list_update_and_delete_linked_visitor_record(self):
        self.client.force_login(self.business)
        created = self.client.post(
            reverse('visitor-records'),
            data=json.dumps(self.visitor_payload()),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 201)
        record_id = created.json()['record']['id']
        self.assertEqual(created.json()['record']['customer']['id'], self.customer.id)
        self.assertEqual(created.json()['record']['visitorCount'], 3)
        self.assertEqual(created.json()['record']['visitorContacts'][1]['role'], '项目总工')
        self.assertEqual(created.json()['record']['hostIds'], [self.business.id, self.other.id])
        self.assertEqual(created.json()['record']['remarks'], '客户自驾到访')
        self.assertEqual(created.json()['record']['customer']['source'], '抖音')

        listed = self.client.get(reverse('visitor-records')).json()['records']
        self.assertEqual(len(listed), 1)
        self.assertTrue(listed[0]['canManage'])

        payload = self.visitor_payload()
        payload['purpose'] = '技术方案交流'
        updated = self.client.patch(
            reverse('visitor-record-detail', args=(record_id,)),
            data=json.dumps(payload),
            content_type='application/json',
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json()['record']['purpose'], '技术方案交流')
        self.assertEqual(self.client.delete(reverse('visitor-record-detail', args=(record_id,))).status_code, 200)
        self.assertFalse(VisitorRecord.objects.filter(id=record_id).exists())

    def test_import_creates_linked_records_with_multiple_people(self):
        self.client.force_login(self.business)
        response = self.client.post(
            reverse('visitor-records-import'),
            data=json.dumps({'records': [{
                '序号': 1,
                '考察时间': '2026年8月13日',
                '关联的客资': self.customer.name,
                '来访客户负责人': '张伟｜项目经理；李敏｜项目总工',
                '会谈内容': '技术方案交流',
                '备注': '客户自驾到访',
                '客户来源（自媒体）': '抖音',
                '接待人员': f'{self.business.username}、{self.other.username}',
                '洽谈总结': '客户认可方案，准备报价',
            }]}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['createdCount'], 1)
        record = VisitorRecord.objects.get()
        self.assertEqual(record.customer, self.customer)
        self.assertEqual(record.host, self.business)
        self.assertEqual(record.visitor_contacts[1], {'name': '李敏', 'role': '项目总工'})
        self.assertEqual(set(record.hosts.values_list('id', flat=True)), {self.business.id, self.other.id})
        self.assertEqual(record.purpose, '技术方案交流')
        self.assertEqual(record.remarks, '客户自驾到访\n客户认可方案，准备报价')
        self.assertEqual(record.notes, '')

    def test_import_is_atomic_when_a_customer_cannot_be_linked(self):
        self.client.force_login(self.business)
        base_row = {
            '序号': 1,
            '考察时间': '2026-08-13',
            '关联的客资': self.customer.name,
            '会谈内容': '方案交流',
            '备注': '',
            '客户来源（自媒体）': '抖音',
            '接待人员': self.business.username,
            '洽谈总结': '',
        }
        invalid_row = {**base_row, '序号': 2, '关联的客资': '不存在的客资'}
        response = self.client.post(
            reverse('visitor-records-import'),
            data=json.dumps({'records': [base_row, invalid_row]}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('第 3 行', response.json()['error'])
        self.assertFalse(VisitorRecord.objects.exists())

    def test_export_returns_multiple_people_and_exact_eight_columns(self):
        record = VisitorRecord.objects.create(
            customer=self.customer,
            visit_date=date(2026, 8, 13),
            visitor_count=2,
            visitor_contacts=[
                {'name': '张伟', 'role': '项目经理'},
                {'name': '李敏', 'role': '项目总工'},
            ],
            purpose='技术方案交流',
            remarks='客户自驾到访',
            notes='客户认可方案，准备报价',
            host=self.business,
            created_by=self.business,
        )
        record.hosts.set([self.business, self.other])
        self.client.force_login(self.business)
        response = self.client.get(reverse('visitor-records-export'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment;', response['Content-Disposition'])
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(
            [sheet.cell(row=1, column=index).value for index in range(1, 9)],
            ['序号', '考察时间', '关联的客资', '来访客户负责人', '会谈内容', '备注', '客户来源（自媒体）', '接待人员'],
        )
        self.assertEqual(sheet.max_column, 8)
        self.assertEqual(sheet['A2'].value, 1)
        self.assertEqual(sheet['B2'].value.date(), date(2026, 8, 13))
        self.assertEqual(sheet['C2'].value, self.customer.name)
        self.assertEqual(sheet['D2'].value, '张伟（项目经理）\n李敏（项目总工）')
        self.assertEqual(sheet['E2'].value, '技术方案交流')
        self.assertEqual(sheet['F2'].value, '客户自驾到访\n客户认可方案，准备报价')
        self.assertEqual(sheet['G2'].value, '抖音')
        self.assertEqual(sheet['H2'].value, f'{self.business.username}\n{self.other.username}')

    def test_legacy_summary_is_visible_as_remarks_and_consolidated_on_update(self):
        record = VisitorRecord.objects.create(
            customer=self.customer,
            visit_date=date(2026, 8, 13),
            visitor_count=1,
            purpose='旧记录',
            remarks='原备注',
            notes='旧洽谈总结',
            host=self.business,
            created_by=self.business,
        )
        self.client.force_login(self.business)

        listed = self.client.get(reverse('visitor-records')).json()['records']
        self.assertEqual(listed[0]['remarks'], '原备注\n旧洽谈总结')
        self.assertNotIn('notes', listed[0])

        payload = self.visitor_payload()
        payload['remarks'] = listed[0]['remarks']
        updated = self.client.patch(
            reverse('visitor-record-detail', args=(record.id,)),
            data=json.dumps(payload),
            content_type='application/json',
        )
        self.assertEqual(updated.status_code, 200)
        record.refresh_from_db()
        self.assertEqual(record.remarks, '原备注\n旧洽谈总结')
        self.assertEqual(record.notes, '')

    def test_filters_by_date_source_keyword_and_any_host(self):
        record = VisitorRecord.objects.create(
            customer=self.customer,
            visit_date=date(2026, 8, 13),
            visitor_count=2,
            contact_name='张伟、李敏',
            visitor_contacts=[{'name': '张伟', 'role': '项目经理'}, {'name': '李敏', 'role': '项目总工'}],
            purpose='技术交流',
            host=self.business,
            created_by=self.business,
        )
        record.hosts.set([self.business, self.other])
        self.client.force_login(self.business)
        response = self.client.get(reverse('visitor-records'), {
            'q': '项目总工',
            'dateFrom': '2026-08-01',
            'dateTo': '2026-08-31',
            'source': '抖音',
            'hostId': self.other.id,
        })
        self.assertEqual([item['id'] for item in response.json()['records']], [record.id])

    def test_unrelated_non_business_user_cannot_change_record(self):
        record = VisitorRecord.objects.create(
            customer=self.customer,
            visit_date=timezone.localdate(),
            visitor_count=1,
            purpose='来访',
            host=self.business,
            created_by=self.business,
        )
        self.client.force_login(self.other)
        response = self.client.delete(reverse('visitor-record-detail', args=(record.id,)))
        self.assertEqual(response.status_code, 403)
        self.assertTrue(VisitorRecord.objects.filter(id=record.id).exists())

    def test_business_level_is_returned_for_group_sorting(self):
        self.client.force_login(self.business)
        employees = self.client.get(reverse('available-employees')).json()
        business = next(item for item in employees['business'] if item['id'] == self.business.id)
        self.assertEqual(business['businessLevel'], 4)


class AvailableEmployeesApiTests(TestCase):
    def create_employee(self, username, role, is_active=True):
        user = get_user_model().objects.create_user(
            username=username,
            password='test-password',
            first_name=username.title(),
            is_active=is_active,
        )
        EmployeeProfile.objects.create(user=user, role=role)
        return user

    def test_returns_only_active_accounts_for_matching_roles(self):
        business = self.create_employee('business', EmployeeProfile.Role.BUSINESS)
        technical = self.create_employee('technical', EmployeeProfile.Role.TECHNICAL)
        both = self.create_employee('both', EmployeeProfile.Role.BOTH)
        self.create_employee('inactive', EmployeeProfile.Role.BOTH, is_active=False)
        get_user_model().objects.create_user(username='unconfigured', password='test-password')

        self.client.force_login(business)
        response = self.client.get(reverse('available-employees'))
        self.assertEqual(response.status_code, 200)
        payload = response.json()

        self.assertEqual(
            {item['id'] for item in payload['business']},
            {business.id, both.id},
        )
        self.assertEqual(
            {item['id'] for item in payload['technical']},
            {technical.id, both.id},
        )
        self.assertEqual(
            {item['id'] for item in payload['all']},
            {business.id, technical.id, both.id},
        )
        business_payload = next(item for item in payload['all'] if item['id'] == business.id)
        technical_payload = next(item for item in payload['all'] if item['id'] == technical.id)
        both_payload = next(item for item in payload['all'] if item['id'] == both.id)
        self.assertIsNotNone(business_payload['businessLevel'])
        self.assertIsNone(business_payload['technicalLevel'])
        self.assertIsNone(technical_payload['businessLevel'])
        self.assertIsNotNone(technical_payload['technicalLevel'])
        self.assertIsNotNone(both_payload['businessLevel'])
        self.assertIsNotNone(both_payload['technicalLevel'])

    def test_admin_is_not_a_business_employee_or_assignable_owner(self):
        superuser = get_user_model().objects.create_superuser(
            username='admin',
            password='test-password',
        )

        self.client.force_login(superuser)
        payload = self.client.get(reverse('available-employees')).json()

        for role in ('all', 'business', 'technical'):
            self.assertNotIn(superuser.id, {item['id'] for item in payload[role]})
        from .views import get_available_owner
        self.assertIsNone(get_available_owner(superuser.id, (EmployeeProfile.Role.BOTH,)))
        followups = self.client.get(reverse('follow-up-tasks')).json()
        self.assertNotIn(superuser.id, {item['id'] for item in followups['employees']})

    def test_owner_choices_sort_by_the_relevant_level_descending(self):
        low = self.create_employee('a-low', EmployeeProfile.Role.BOTH)
        business = self.create_employee('b-business', EmployeeProfile.Role.BOTH)
        technical = self.create_employee('c-technical', EmployeeProfile.Role.BOTH)
        tied = self.create_employee('d-tied', EmployeeProfile.Role.BOTH)
        EmployeeProfile.objects.filter(user=business).update(business_level=5, technical_level=2)
        EmployeeProfile.objects.filter(user=technical).update(business_level=2, technical_level=5)
        EmployeeProfile.objects.filter(user=tied).update(business_level=5, technical_level=5)
        self.client.force_login(low)
        payload = self.client.get(reverse('available-employees')).json()
        self.assertEqual([item['id'] for item in payload['business']], [business.id, tied.id, technical.id, low.id])
        self.assertEqual([item['id'] for item in payload['technical']], [technical.id, tied.id, business.id, low.id])

    def test_anonymous_request_is_redirected_to_login(self):
        response = self.client.get(reverse('available-employees'))
        self.assertRedirects(
            response,
            f"{reverse('login')}?next={reverse('available-employees')}",
        )


class LoginAndAdminFormTests(TestCase):
    def test_employee_profile_admin_loads_role_based_level_visibility_script(self):
        admin_user = get_user_model().objects.create_superuser(
            username='profile-admin',
            password='123456',
        )
        employee = get_user_model().objects.create_user(username='tech-profile')
        profile = EmployeeProfile.objects.create(
            user=employee,
            role=EmployeeProfile.Role.TECHNICAL,
            technical_level=3,
        )
        self.client.force_login(admin_user)

        response = self.client.get(reverse('admin:customers_employeeprofile_change', args=(profile.id,)))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="business_level"')
        self.assertContains(response, 'name="technical_level"')
        self.assertContains(response, 'admin/employee-profile-levels.js')

    def test_system_management_link_is_only_visible_to_superusers(self):
        employee = get_user_model().objects.create_user(
            username='employee',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=employee,
            role=EmployeeProfile.Role.BUSINESS,
        )
        self.client.force_login(employee)
        response = self.client.get(reverse('home'))
        self.assertNotContains(response, 'href="/admin/"')

        administrator = get_user_model().objects.create_user(
            username='administrator',
            password='123456',
            is_staff=True,
        )
        EmployeeProfile.objects.create(
            user=administrator,
            role=EmployeeProfile.Role.BOTH,
        )
        self.client.force_login(administrator)
        response = self.client.get(reverse('home'))
        self.assertNotContains(response, 'href="/admin/"')

        superuser = get_user_model().objects.create_superuser(
            username='superuser',
            password='123456',
        )
        self.client.force_login(superuser)
        response = self.client.get(reverse('home'))
        self.assertContains(response, 'href="/admin/"')

    def test_home_requires_login_and_six_digit_numeric_password_can_login(self):
        user = get_user_model().objects.create_user(
            username='employee',
            password='123456',
        )
        response = self.client.get(reverse('home'))
        self.assertRedirects(response, f"{reverse('login')}?next={reverse('home')}")

        response = self.client.post(
            reverse('login'),
            {'username': user.username, 'password': '123456'},
        )
        self.assertRedirects(response, reverse('home'))

    def test_admin_add_user_form_only_contains_username_and_passwords(self):
        admin_user = get_user_model().objects.create_superuser(
            username='admin',
            password='123456',
        )
        self.client.force_login(admin_user)

        response = self.client.get(reverse('admin:auth_user_add'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="username"')
        self.assertContains(response, 'name="password1"')
        self.assertContains(response, 'name="password2"')
        self.assertContains(response, 'name="is_staff"')
        self.assertNotContains(response, 'name="usable_password"')
        self.assertNotContains(response, 'name="first_name"')
        self.assertNotContains(response, 'name="last_name"')
        self.assertNotContains(response, 'name="email"')

    def test_new_employee_is_staff_by_default_and_can_be_unchecked(self):
        form = EmployeeUserCreationForm(
            data={
                'username': 'default-staff',
                'password1': '123456',
                'password2': '123456',
                'is_staff': 'on',
            },
        )
        self.assertTrue(form.is_valid(), form.errors)
        user = form.save()
        self.assertTrue(user.is_staff)

        form = EmployeeUserCreationForm(
            data={
                'username': 'non-staff',
                'password1': '123456',
                'password2': '123456',
            },
        )
        self.assertTrue(form.is_valid(), form.errors)
        user = form.save()
        self.assertFalse(user.is_staff)

    def test_admin_change_page_hides_password_hash_and_only_shows_reset_action(self):
        admin_user = get_user_model().objects.create_superuser(
            username='admin',
            password='123456',
        )
        employee = get_user_model().objects.create_user(
            username='employee',
            password='654321',
            is_staff=True,
        )
        self.client.force_login(admin_user)

        response = self.client.get(
            reverse('admin:auth_user_change', args=(employee.pk,)),
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '密码已安全加密，原始密码无法查看')
        self.assertContains(response, '重置密码')
        self.assertContains(
            response,
            reverse('admin:auth_user_password_change', args=(employee.pk,)),
        )
        self.assertNotContains(response, 'pbkdf2_sha256')
        self.assertNotContains(response, '迭代次数')
        self.assertNotContains(response, '盐：')


class CustomerPersistenceApiTests(TestCase):
    def setUp(self):
        self.business = get_user_model().objects.create_user(
            username='business-owner',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.business,
            role=EmployeeProfile.Role.BUSINESS,
            nickname='商务甲',
        )
        self.technical = get_user_model().objects.create_user(
            username='technical-owner',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.technical,
            role=EmployeeProfile.Role.TECHNICAL,
            nickname='技术乙',
        )
        self.customer = Customer.objects.create(
            name='修改前名称',
            phone='13800000000',
            grade='B',
            province='浙江省',
            city='杭州市',
            district='西湖区',
            business_owner=self.business,
            technical_owner=self.technical,
        )
        self.client.force_login(self.business)

    def test_customer_edit_is_persisted_and_returned_by_list_api(self):
        payload = {
            'name': '修改后名称',
            'phone': '13800000000',
            'source': '朋友介绍',
            'referrer': '王先生',
            'province': '广东省',
            'city': '佛山市',
            'district': '顺德区',
            'grade': 'A',
            'cooperationStatus': 'cooperating',
            'description': '数据库持久化测试',
            'progress': '方案与报价',
            'plan': '修改后的施工方案',
            'ownerId': self.business.id,
            'techId': self.technical.id,
        }
        response = self.client.patch(
            reverse('customer-detail', args=[self.customer.id]),
            data=json.dumps(payload),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.name, '修改后名称')
        self.assertEqual(self.customer.district, '顺德区')
        self.assertEqual(self.customer.referrer, '王先生')
        self.assertEqual(self.customer.cooperation_status, Customer.CooperationStatus.COOPERATING)

        response = self.client.get(reverse('customers-collection'))
        payload = response.json()
        saved = next(
            item for item in payload['customers']
            if item['id'] == self.customer.id
        )
        self.assertEqual(saved['name'], '修改后名称')
        self.assertEqual(saved['description'], '数据库持久化测试')
        self.assertEqual(saved['cooperationStatus'], 'cooperating')
        self.assertEqual(saved['cooperationStatusLabel'], '已建立合作关系')
        self.assertEqual(saved['owner'], '商务甲')
        self.assertEqual(saved['tech'], '技术乙')
        self.assertRegex(saved['createdDate'], r'^\d{4}-\d{2}-\d{2}$')
        self.assertTrue(
            any(project['customer']['id'] == self.customer.id for project in payload['projects']),
        )

    def test_customer_edit_accepts_wechat_contact(self):
        response = self.client.patch(
            reverse('customer-detail', args=[self.customer.id]),
            data=json.dumps({
                'name': self.customer.name,
                'phone': '微信：wx_customer_01',
                'source': '抖音',
                'province': self.customer.province,
                'city': self.customer.city,
                'district': self.customer.district,
                'grade': self.customer.grade,
                'cooperationStatus': 'none',
                'description': '',
                'progress': '需求对接',
                'plan': '原方案',
                'ownerId': self.business.id,
                'techId': self.technical.id,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.phone, '微信：wx_customer_01')

    def test_service_account_source_and_wechat_status_are_persisted(self):
        response = self.client.patch(
            reverse('customer-detail', args=[self.customer.id]),
            data=json.dumps({
                'name': self.customer.name,
                'phone': self.customer.phone,
                'wechatStatus': 'yes',
                'source': '服务号',
                'province': self.customer.province,
                'city': self.customer.city,
                'district': self.customer.district,
                'grade': self.customer.grade,
                'cooperationStatus': 'none',
                'description': '',
                'progress': '需求对接',
                'plan': '原方案',
                'ownerId': self.business.id,
                'techId': self.technical.id,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.source, Customer.Source.SERVICE_ACCOUNT)
        self.assertEqual(self.customer.wechat_status, 'yes')
        self.assertEqual(response.json()['customer']['source'], '服务号')
        self.assertEqual(response.json()['customer']['wechatStatus'], 'yes')

    def test_unsupported_customer_source_is_rejected_without_changing_data(self):
        response = self.client.patch(
            reverse('customer-detail', args=[self.customer.id]),
            data=json.dumps({
                'name': self.customer.name,
                'phone': self.customer.phone,
                'source': '其他',
                'province': self.customer.province,
                'city': self.customer.city,
                'district': self.customer.district,
                'grade': self.customer.grade,
                'cooperationStatus': 'none',
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('抖音、视频号、服务号、朋友介绍', response.json()['error'])
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.source, Customer.Source.DOUYIN)

    def test_legacy_source_can_be_preserved_during_an_unrelated_edit(self):
        self.customer.source = '其他'
        self.customer.save(update_fields=('source',))
        response = self.client.patch(
            reverse('customer-detail', args=[self.customer.id]),
            data=json.dumps({
                'name': '历史来源名称已修改',
                'phone': self.customer.phone,
                'source': '其他',
                'province': self.customer.province,
                'city': self.customer.city,
                'district': self.customer.district,
                'grade': self.customer.grade,
                'cooperationStatus': 'none',
                'ownerId': self.business.id,
                'techId': self.technical.id,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.name, '历史来源名称已修改')
        self.assertEqual(self.customer.source, '其他')

    def test_batch_lead_creation_accepts_qq_contact(self):
        response = self.client.post(
            reverse('customers-import'),
            data=json.dumps({
                'mode': 'leads',
                'customers': [{
                    'name': 'QQ联系方式客资',
                    'phone': 'QQ：1234567',
                    'description': '',
                    'source': '抖音',
                }],
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(Customer.objects.get(name='QQ联系方式客资').phone, 'QQ：1234567')

    def test_batch_lead_creation_accepts_service_account_source(self):
        response = self.client.post(
            reverse('customers-import'),
            data=json.dumps({
                'mode': 'leads',
                'customers': [{
                    'name': '服务号批量客资',
                    'phone': 'wx_service_account',
                    'description': '来自服务号',
                    'source': '服务号',
                }],
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(Customer.objects.get(name='服务号批量客资').source, Customer.Source.SERVICE_ACCOUNT)

    def test_new_customer_has_project_but_no_automatic_partnership_identity(self):
        response = self.client.post(
            reverse('customers-collection'),
            data=json.dumps({
                'name': '无合作身份新客户',
                'phone': '13800000009',
                'source': '抖音',
                'province': '浙江省',
                'city': '杭州市',
                'district': '余杭区',
                'grade': 'B',
                'projectName': '无合作身份客户首个项目',
                'progress': '需求对接',
                'description': '客户资料与合作身份相互独立',
                'plan': '',
                'projectTypeId': '',
                'ownerId': self.business.id,
                'techId': self.technical.id,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 201)
        customer = Customer.objects.get(id=response.json()['customer']['id'])
        self.assertTrue(customer.projects.filter(is_active=True).exists())
        self.assertEqual(customer.projects.get().name, '无合作身份客户首个项目')
        self.assertEqual(response.json()['customer']['projectName'], '无合作身份客户首个项目')
        self.assertFalse(customer.partnership_identities.filter(is_active=True).exists())
        self.assertEqual(customer.cooperation_status, Customer.CooperationStatus.NONE)
        self.assertEqual(response.json()['customer']['cooperationStatus'], 'none')

    def test_c_grade_only_saves_basic_and_location_fields_with_optional_district(self):
        project_type = ProjectType.objects.create(name='不应写入的类型')
        response = self.client.post(
            reverse('customers-collection'),
            data=json.dumps({
                'name': 'C级精简客资',
                'phone': '13800000031',
                'source': '抖音',
                'province': '河北省',
                'city': '唐山市',
                'district': '',
                'grade': 'C',
                'projectName': 'C级项目名称',
                'progress': '合同签订',
                'description': '只登记前两栏',
                'plan': '不应写入的施工方案',
                'projectTypeId': project_type.id,
                'ownerId': self.business.id,
                'techId': self.technical.id,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 201)
        customer = Customer.objects.get(id=response.json()['customer']['id'])
        project = customer.projects.get()
        self.assertEqual(customer.district, '')
        self.assertEqual(customer.progress, '需求对接')
        self.assertEqual(customer.plan, '')
        self.assertIsNone(customer.business_owner)
        self.assertIsNone(customer.technical_owner)
        self.assertEqual(project.name, 'C级项目名称')
        self.assertEqual(project.progress, '需求对接')
        self.assertEqual(project.plan, '')
        self.assertIsNone(project.project_type)
        self.assertIsNone(project.business_owner)
        self.assertIsNone(project.technical_owner)

    def test_new_customer_district_is_optional_for_all_grades(self):
        for index, grade in enumerate(('A', 'B', 'C', 'D')):
            with self.subTest(grade=grade):
                payload = {
                    'name': f'{grade}级区县选填客资',
                    'phone': f'1380000004{index}',
                    'grade': grade,
                    'province': '河北省',
                    'city': '唐山市',
                    'projectName': f'{grade}级项目',
                }
                if grade in ('A', 'B'):
                    payload.update(ownerId=self.business.id, techId=self.technical.id)
                response = self.client.post(
                    reverse('customers-collection'),
                    data=json.dumps(payload),
                    content_type='application/json',
                )
                self.assertEqual(response.status_code, 201)
                customer = Customer.objects.get(id=response.json()['customer']['id'])
                self.assertEqual(customer.district, '')
                if grade in ('C', 'D'):
                    self.assertIsNone(customer.business_owner)
                    self.assertIsNone(customer.technical_owner)

    def test_d_grade_edit_ignores_hidden_fields_without_erasing_project_history(self):
        self.customer.progress = '方案与报价'
        self.customer.plan = '保留的历史方案'
        self.customer.save(update_fields=('progress', 'plan', 'updated_at'))
        project = Project.objects.create(
            customer=self.customer,
            name='降级前项目',
            progress='方案与报价',
            plan='保留的历史方案',
            province='浙江省',
            city='杭州市',
            business_owner=self.business,
            technical_owner=self.technical,
            owners_inherit_customer=False,
        )

        response = self.client.patch(
            reverse('customer-detail', args=[self.customer.id]),
            data=json.dumps({
                'name': self.customer.name,
                'phone': self.customer.phone,
                'source': '抖音',
                'province': '浙江省',
                'city': '杭州市',
                'district': '',
                'grade': 'D',
                'projectName': project.name,
                'progress': '合同签订',
                'plan': '不应覆盖的隐藏方案',
                'projectTypeId': '',
                'ownerId': 'unavailable-hidden-owner',
                'techId': 'unavailable-hidden-technician',
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        self.customer.refresh_from_db()
        project.refresh_from_db()
        self.assertEqual(self.customer.grade, 'D')
        self.assertEqual(self.customer.district, '')
        self.assertEqual(self.customer.progress, '方案与报价')
        self.assertEqual(self.customer.plan, '保留的历史方案')
        self.assertEqual(self.customer.business_owner, self.business)
        self.assertEqual(self.customer.technical_owner, self.technical)
        self.assertEqual(project.progress, '方案与报价')
        self.assertEqual(project.plan, '保留的历史方案')
        self.assertEqual(project.business_owner, self.business)
        self.assertEqual(project.technical_owner, self.technical)
        self.assertFalse(project.owners_inherit_customer)

    def test_new_customer_requires_project_name_without_creating_partial_customer(self):
        original_count = Customer.objects.count()
        response = self.client.post(
            reverse('customers-collection'),
            data=json.dumps({
                'name': '缺少项目名的客户',
                'phone': '13800000010',
                'source': '抖音',
                'province': '浙江省',
                'city': '杭州市',
                'district': '余杭区',
                'grade': 'B',
                'projectName': '   ',
                'progress': '需求对接',
                'projectTypeId': '',
                'ownerId': self.business.id,
                'techId': self.technical.id,
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()['error'], '请填写项目名称')
        self.assertEqual(Customer.objects.count(), original_count)

    def test_legacy_customer_can_be_edited_without_filling_unrelated_missing_fields(self):
        legacy = Customer.objects.create(
            name='旧客资',
            phone='13800000022',
            province='浙江省',
            city='杭州市',
            district='',
            business_owner_name='原商务负责人',
            technical_owner_name='原技术负责人',
        )
        payload = {
            'name': '旧客资已修改',
            'phone': '13800000022',
            'source': '抖音',
            'referrer': '',
            'province': '浙江省',
            'city': '杭州市',
            'district': '',
            'grade': 'B',
            'description': '只修改这段客户描述',
            'progress': '方案与报价',
            'plan': '原方案',
            'ownerId': '',
            'techId': '',
        }
        response = self.client.patch(
            reverse('customer-detail', args=[legacy.id]),
            data=json.dumps(payload),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        legacy.refresh_from_db()
        self.assertEqual(legacy.description, '只修改这段客户描述')
        self.assertEqual(legacy.district, '')
        self.assertEqual(legacy.business_owner_name, '原商务负责人')
        self.assertEqual(legacy.technical_owner_name, '原技术负责人')

    def test_customer_import_appends_without_replacing_existing_records(self):
        existing_id = self.customer.id
        original_count = Customer.objects.count()
        payload = {
            'customers': [
                {
                    'name': '追加导入甲',
                    'phone': '13800000001',
                    'source': '抖音',
                    'province': '河北省',
                    'city': '石家庄市',
                    'district': '长安区',
                    'grade': 'A',
                    'description': '第一条追加导入测试',
                    'projectName': '河北道路一期',
                    'progress': '需求对接',
                    'plan': '道路方案',
                    'projectTypeId': '',
                },
                {
                    'name': '追加导入乙',
                    'phone': '13800000002',
                    'source': '视频号',
                    'province': '山东省',
                    'city': '济南市',
                    'district': '历下区',
                    'grade': 'B',
                    'description': '第二条追加导入测试',
                    'projectName': '济南回填一期',
                    'progress': '需求对接',
                    'plan': '回填方案',
                    'projectTypeId': '',
                },
            ],
        }

        response = self.client.post(
            reverse('customers-import'),
            data=json.dumps(payload),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['createdCount'], 2)
        self.assertEqual(Customer.objects.count(), original_count + 2)
        self.assertTrue(Customer.objects.filter(id=existing_id, name='修改前名称').exists())
        self.assertTrue(Customer.objects.filter(name='追加导入甲').exists())
        self.assertTrue(Customer.objects.filter(name='追加导入乙').exists())
        self.assertEqual(Project.objects.get(customer__name='追加导入甲').name, '河北道路一期')
        self.assertEqual(Project.objects.get(customer__name='追加导入乙').name, '济南回填一期')

    def test_batch_lead_creation_only_saves_the_four_initial_fields(self):
        existing_id = self.customer.id
        payload = {
            'mode': 'leads',
            'customers': [
                {
                    'name': '批量线索甲',
                    'phone': '13800000011',
                    'description': '先登记需求，后续补充项目信息',
                    'source': '视频号',
                },
                {
                    'name': '批量线索乙',
                    'phone': '13800000014',
                    'description': '',
                    'source': '抖音',
                },
            ],
        }

        response = self.client.post(
            reverse('customers-import'),
            data=json.dumps(payload),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['createdCount'], 2)
        self.assertTrue(Customer.objects.filter(id=existing_id).exists())
        first = Customer.objects.get(name='批量线索甲')
        second = Customer.objects.get(name='批量线索乙')
        self.assertEqual(first.phone, '13800000011')
        self.assertEqual(first.description, '先登记需求，后续补充项目信息')
        self.assertEqual(first.source, '视频号')
        self.assertEqual(first.channel, '批量新增')
        self.assertEqual(first.grade, 'C')
        self.assertEqual(second.grade, 'C')
        self.assertEqual(first.province, '')
        self.assertEqual(first.city, '')
        self.assertIsNone(first.business_owner)
        self.assertIsNone(first.technical_owner)
        self.assertFalse(first.projects.exists())
        self.assertFalse(second.projects.exists())

        partial_update = self.client.patch(
            reverse('customer-detail', args=[first.id]),
            data=json.dumps({
                'name': '批量线索甲',
                'phone': '13800000013',
                'description': '第二次沟通，仍待确认项目',
                'source': '视频号',
                'grade': 'B',
                'cooperationStatus': 'none',
                'progress': '需求对接',
                'projectTypeId': '',
                'projectName': '',
                'allowIncomplete': True,
            }),
            content_type='application/json',
        )
        self.assertEqual(partial_update.status_code, 200)
        first.refresh_from_db()
        self.assertEqual(first.phone, '13800000013')
        self.assertEqual(first.description, '第二次沟通，仍待确认项目')
        self.assertFalse(first.projects.exists())

        project_update = self.client.patch(
            reverse('customer-detail', args=[first.id]),
            data=json.dumps({
                'name': '批量线索甲',
                'phone': '13800000013',
                'description': '项目已确认',
                'source': '视频号',
                'province': '河北省',
                'city': '石家庄市',
                'district': '',
                'grade': 'B',
                'cooperationStatus': 'none',
                'progress': '需求对接',
                'projectTypeId': '',
                'projectName': '石家庄道路项目',
                'allowIncomplete': True,
            }),
            content_type='application/json',
        )
        self.assertEqual(project_update.status_code, 200)
        self.assertEqual(project_update.json()['project']['name'], '石家庄道路项目')
        self.assertEqual(first.projects.get().name, '石家庄道路项目')

    def test_batch_lead_creation_is_atomic_when_a_filled_row_has_no_name(self):
        original_ids = list(Customer.objects.values_list('id', flat=True))
        response = self.client.post(
            reverse('customers-import'),
            data=json.dumps({
                'mode': 'leads',
                'customers': [
                    {'name': '本行原本有效', 'phone': '13800000015', 'description': '', 'source': '抖音'},
                    {'name': '', 'phone': '13800000012', 'description': '', 'source': '视频号'},
                ],
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('第2行：请填写客户名称', response.json()['error'])
        self.assertEqual(list(Customer.objects.values_list('id', flat=True)), original_ids)

    def test_customer_import_is_atomic_when_one_row_is_invalid(self):
        original_ids = list(Customer.objects.values_list('id', flat=True))
        payload = {
            'customers': [
                {
                    'name': '本行原本有效',
                    'phone': '13800000016',
                    'source': '抖音',
                    'province': '河北省',
                    'city': '石家庄市',
                    'grade': 'A',
                    'projectName': '本行有效项目',
                    'progress': '需求对接',
                    'projectTypeId': '',
                },
                {
                    'name': '缺少城市的无效行',
                    'phone': '13800000017',
                    'source': '视频号',
                    'province': '山东省',
                    'city': '',
                    'grade': 'B',
                    'projectName': '缺少城市项目',
                    'progress': '需求对接',
                    'projectTypeId': '',
                },
            ],
        }

        response = self.client.post(
            reverse('customers-import'),
            data=json.dumps(payload),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('第3行', response.json()['error'])
        self.assertEqual(list(Customer.objects.values_list('id', flat=True)), original_ids)

    def test_technical_employee_cannot_delete_customer(self):
        self.client.force_login(self.technical)
        response = self.client.delete(
            reverse('customer-detail', args=[self.customer.id]),
        )
        self.assertEqual(response.status_code, 403)
        self.assertTrue(Customer.objects.filter(id=self.customer.id).exists())

    def test_customer_owner_changes_are_synced_to_every_project_task(self):
        project = Project.objects.create(
            customer=self.customer,
            name='负责人同步测试项目',
            progress='需求对接',
        )
        seed_historical_stage_tasks(project)
        business_task = project.follow_up_tasks.get(
            role=FollowUpTask.Role.BUSINESS,
        )
        technical_task = FollowUpTask.objects.create(
            project=project,
            title='技术临时任务',
            target_progress='项目内临时任务',
            role=FollowUpTask.Role.TECHNICAL,
            is_manual=True,
        )

        payload = {
            'name': self.customer.name,
            'phone': self.customer.phone,
            'source': self.customer.source,
            'referrer': self.customer.referrer,
            'province': self.customer.province,
            'city': self.customer.city,
            'district': self.customer.district,
            'grade': self.customer.grade,
            'description': self.customer.description,
            'progress': '需求对接',
            'plan': '负责人同步测试方案',
            'ownerId': self.business.id,
            'techId': self.technical.id,
        }
        response = self.client.patch(
            reverse('customer-detail', args=[self.customer.id]),
            data=json.dumps(payload),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        project.refresh_from_db()
        business_task.refresh_from_db()
        technical_task.refresh_from_db()
        self.assertEqual(project.business_owner, self.business)
        self.assertEqual(project.technical_owner, self.technical)
        self.assertEqual(business_task.assignee, self.business)
        self.assertEqual(technical_task.assignee, self.technical)

    def test_project_specific_owner_is_not_overwritten_by_customer_default(self):
        independent_owner = get_user_model().objects.create_user(
            username='independent-project-owner', password='123456',
        )
        EmployeeProfile.objects.create(
            user=independent_owner,
            role=EmployeeProfile.Role.BUSINESS,
            nickname='项目独立负责人',
        )
        project = Project.objects.create(
            customer=self.customer,
            name='独立负责人项目',
            progress='需求对接',
            business_owner=independent_owner,
            technical_owner=self.technical,
            owners_inherit_customer=False,
        )
        response = self.client.patch(
            reverse('customer-detail', args=[self.customer.id]),
            data=json.dumps({
                'name': self.customer.name,
                'phone': self.customer.phone,
                'source': self.customer.source,
                'province': self.customer.province,
                'city': self.customer.city,
                'district': self.customer.district,
                'grade': self.customer.grade,
                'description': self.customer.description,
                'progress': '需求对接',
                'plan': '客户默认方案',
                'ownerId': self.business.id,
                'techId': self.technical.id,
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        project.refresh_from_db()
        self.assertEqual(project.business_owner, independent_owner)
        self.assertFalse(project.owners_inherit_customer)

    def test_follow_up_owner_uses_matching_legacy_department_name(self):
        legacy = Customer.objects.create(
            name='旧数据负责人展示测试',
            grade='B',
            business_owner_name='旧商务负责人',
            technical_owner_name='旧技术负责人',
        )
        project = Project.objects.create(
            customer=legacy,
            name='旧数据项目',
            progress='需求对接',
        )
        seed_historical_stage_tasks(project)
        FollowUpTask.objects.create(
            project=project,
            title='技术临时任务',
            target_progress='项目内临时任务',
            role=FollowUpTask.Role.TECHNICAL,
            is_manual=True,
        )

        response = self.client.get(
            reverse('customer-detail', args=[legacy.id]),
        )

        self.assertEqual(response.status_code, 200)
        owners = {
            task['role']: task['assignee']['name']
            for task in response.json()['followUps']
        }
        self.assertEqual(owners[FollowUpTask.Role.BUSINESS], '旧商务负责人')
        self.assertEqual(owners[FollowUpTask.Role.TECHNICAL], '旧技术负责人')

    def test_business_employee_can_delete_customer_and_related_records(self):
        project = Project.objects.create(
            customer=self.customer,
            name='删除测试项目',
            progress='需求对接',
        )
        seed_historical_stage_tasks(project)
        task = project.follow_up_tasks.get(
            target_progress='技术验证',
            role=FollowUpTask.Role.BUSINESS,
        )
        response = self.client.delete(
            reverse('customer-detail', args=[self.customer.id]),
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['deleted'])
        self.assertEqual(payload['projectCount'], 1)
        self.assertEqual(payload['followUpCount'], 2)
        self.assertFalse(Customer.objects.filter(id=self.customer.id).exists())
        self.assertFalse(Project.objects.filter(id=project.id).exists())
        self.assertFalse(FollowUpTask.objects.filter(id=task.id).exists())

    def test_superuser_without_business_role_can_delete_customer(self):
        admin = get_user_model().objects.create_superuser(
            username='customer-delete-admin',
            password='123456',
            email='admin@example.com',
        )
        self.client.force_login(admin)

        response = self.client.delete(
            reverse('customer-detail', args=[self.customer.id]),
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(Customer.objects.filter(id=self.customer.id).exists())


class FollowUpTaskApiTests(TestCase):
    def setUp(self):
        self.business = get_user_model().objects.create_user(
            username='follow-business',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.business,
            role=EmployeeProfile.Role.BUSINESS,
        )
        self.technical = get_user_model().objects.create_user(
            username='follow-technical',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.technical,
            role=EmployeeProfile.Role.TECHNICAL,
        )
        self.other = get_user_model().objects.create_user(
            username='follow-other',
            password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.other,
            role=EmployeeProfile.Role.BUSINESS,
        )
        self.admin = get_user_model().objects.create_superuser(
            username='follow-admin',
            password='123456',
        )
        self.customer = Customer.objects.create(
            name='自动跟进测试客户',
            phone='13800000031',
            grade='B',
            progress=PROGRESS_STAGES[4],
            business_owner=self.business,
            technical_owner=self.technical,
        )
        self.project = Project.objects.create(
            customer=self.customer,
            name='自动跟进测试项目',
            progress=PROGRESS_STAGES[4],
            business_owner=self.business,
            technical_owner=self.technical,
        )

        seed_historical_stage_tasks(self.project)

    def test_historical_milestone_task_keeps_correct_owner(self):
        task = self.project.follow_up_tasks.get(
            target_progress=PROGRESS_STAGES[5],
        )
        self.assertEqual(task.role, FollowUpTask.Role.TECHNICAL)
        self.assertEqual(task.assignee, self.technical)
        self.assertEqual(task.status, FollowUpTask.Status.PENDING)

    def test_project_progress_uses_the_seven_core_stages(self):
        self.assertEqual(PROGRESS_STAGES, (
            '需求对接',
            '技术验证',
            '客户深度沟通',
            '方案与报价',
            '合同签订',
            '项目实施跟进',
            '售后维护与需求挖掘',
        ))
        self.assertEqual(normalize_progress('B段施工完成'), '项目实施跟进')

    def test_customer_history_hides_system_tasks_beyond_next_stage(self):
        stale = FollowUpTask.objects.create(
            project=self.project,
            title='不应提前显示的回访任务',
            target_progress='售后维护与需求挖掘',
            role=FollowUpTask.Role.BUSINESS,
            assignee=self.business,
            status=FollowUpTask.Status.COMPLETED,
            completed_at=timezone.now(),
            result='旧演示记录',
        )
        self.client.force_login(self.admin)

        response = self.client.get(reverse('customer-detail', args=[self.customer.id]))

        self.assertEqual(response.status_code, 200)
        returned_ids = {task['id'] for task in response.json()['followUps']}
        self.assertNotIn(stale.id, returned_ids)
        self.assertTrue(returned_ids)

    def test_legacy_ab_stage_task_is_displayed_as_unified_trial_stage(self):
        legacy = FollowUpTask.objects.create(
            project=self.project,
            title='推进并完成 B 段施工',
            target_progress='B段施工完成',
            role=FollowUpTask.Role.TECHNICAL,
            assignee=self.technical,
            status=FollowUpTask.Status.COMPLETED,
            completed_at=timezone.now(),
            result='历史施工记录仍保留',
        )
        self.project.progress = '项目实施跟进'
        self.project.save(update_fields=('progress', 'updated_at'))
        self.customer.progress = '项目实施跟进'
        self.customer.save(update_fields=('progress', 'updated_at'))
        self.client.force_login(self.admin)

        detail = self.client.get(reverse('customer-detail', args=[self.customer.id])).json()
        record = next(task for task in detail['followUps'] if task['id'] == legacy.id)
        self.assertEqual(record['targetProgress'], '项目实施跟进')
        self.assertNotIn('A 段', record['title'])
        self.assertNotIn('B 段', record['title'])

    def test_employee_only_sees_tasks_assigned_to_their_account(self):
        self.client.force_login(self.technical)
        response = self.client.get(reverse('follow-up-tasks'))
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertFalse(payload['canViewAll'])
        self.assertTrue(payload['tasks'])
        self.assertEqual(
            {task['assignee']['id'] for task in payload['tasks']},
            {self.technical.id},
        )

        self.client.force_login(self.other)
        response = self.client.get(reverse('follow-up-tasks'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['tasks'], [])

    def test_superuser_can_view_all_and_filter_by_employee(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('follow-up-tasks'))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['canViewAll'])
        self.assertTrue(response.json()['tasks'])

        response = self.client.get(
            reverse('follow-up-tasks'),
            {'employee': self.other.id},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['tasks'], [])

    def test_c_and_d_customers_do_not_generate_or_display_followup_tasks(self):
        hidden_tasks = []
        hidden_customers = []
        for grade in ('C', 'D'):
            customer = Customer.objects.create(
                name=f'{grade}类客户',
                grade=grade,
                progress=PROGRESS_STAGES[4],
                business_owner=self.business,
                technical_owner=self.technical,
            )
            project = Project.objects.create(
                customer=customer,
                name=f'{grade}类项目',
                progress=PROGRESS_STAGES[4],
                business_owner=self.business,
                technical_owner=self.technical,
            )
            self.assertFalse(project.follow_up_tasks.exists())
            hidden_customers.append(customer)
            hidden_tasks.append(FollowUpTask.objects.create(
                project=project,
                title='历史跟进记录',
                target_progress=PROGRESS_STAGES[5],
                role=FollowUpTask.Role.TECHNICAL,
                assignee=self.business,
            ))
        self.client.force_login(self.admin)

        response = self.client.get(reverse('follow-up-tasks'))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        returned_ids = {task['id'] for task in payload['tasks']}
        self.assertTrue(returned_ids.isdisjoint(task.id for task in hidden_tasks))
        self.assertEqual(payload['summary']['total'], len(payload['tasks']))

        for customer in hidden_customers:
            detail = self.client.get(reverse('customer-detail', args=[customer.id]))
            self.assertEqual(detail.status_code, 200)
            self.assertEqual(detail.json()['followUps'], [])

    def test_c_and_d_customers_cannot_create_or_modify_followup_tasks(self):
        self.client.force_login(self.admin)
        for grade in ('C', 'D'):
            customer = Customer.objects.create(
                name=f'{grade}类受限客户',
                grade=grade,
                business_owner=self.business,
                technical_owner=self.technical,
            )
            project = Project.objects.create(
                customer=customer,
                name=f'{grade}类受限项目',
                business_owner=self.business,
                technical_owner=self.technical,
            )
            create_response = self.client.post(
                reverse('follow-up-tasks'),
                data=json.dumps({
                    'projectId': project.id,
                    'title': '不应创建的跟进任务',
                    'role': FollowUpTask.Role.BUSINESS,
                }),
                content_type='application/json',
            )
            self.assertEqual(create_response.status_code, 400)

            historical_task = FollowUpTask.objects.create(
                project=project,
                title='保留但不可操作的历史记录',
                target_progress=PROGRESS_STAGES[1],
                role=FollowUpTask.Role.BUSINESS,
                assignee=self.business,
            )
            detail_url = reverse('follow-up-task-detail', args=[historical_task.id])
            self.assertEqual(
                self.client.patch(
                    detail_url,
                    data=json.dumps({'title': '尝试修改'}),
                    content_type='application/json',
                ).status_code,
                403,
            )
            self.assertEqual(self.client.delete(detail_url).status_code, 403)
            self.assertTrue(FollowUpTask.objects.filter(id=historical_task.id).exists())

    def test_api_returns_only_one_current_status_card_per_project(self):
        current_task = self.project.follow_up_tasks.get(
            target_progress=PROGRESS_STAGES[5],
        )
        FollowUpTask.objects.create(
            project=self.project,
            title='历史阶段任务',
            target_progress=PROGRESS_STAGES[3],
            role=FollowUpTask.Role.BUSINESS,
            status=FollowUpTask.Status.COMPLETED,
            completed_at=timezone.now(),
            result='历史任务已完成',
        )

        self.client.force_login(self.admin)
        response = self.client.get(reverse('follow-up-tasks'))

        self.assertEqual(response.status_code, 200)
        project_tasks = [
            task for task in response.json()['tasks']
            if task['project']['id'] == self.project.id
        ]
        self.assertEqual(len(project_tasks), 1)
        self.assertEqual(project_tasks[0]['id'], current_task.id)
        returned_tasks = response.json()['tasks']
        self.assertEqual(
            len(returned_tasks),
            len({task['project']['id'] for task in returned_tasks}),
        )

    def test_unassigned_employee_cannot_modify_another_employees_task(self):
        task = self.project.follow_up_tasks.get(
            target_progress=PROGRESS_STAGES[5],
        )
        self.client.force_login(self.other)
        response = self.client.patch(
            reverse('follow-up-task-detail', args=[task.id]),
            data=json.dumps({'status': FollowUpTask.Status.WAITING}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 403)

        response = self.client.delete(
            reverse('follow-up-task-detail', args=[task.id]),
        )
        self.assertEqual(response.status_code, 403)
        self.assertTrue(FollowUpTask.objects.filter(id=task.id).exists())

    def test_completion_requires_result_and_does_not_advance_project(self):
        task = self.project.follow_up_tasks.get(
            target_progress=PROGRESS_STAGES[5],
        )
        self.client.force_login(self.technical)
        url = reverse('follow-up-task-detail', args=[task.id])
        response = self.client.patch(
            url,
            data=json.dumps({'status': FollowUpTask.Status.COMPLETED}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

        response = self.client.patch(
            url,
            data=json.dumps({
                'status': FollowUpTask.Status.COMPLETED,
                'result': '已完成现场工作并上传记录',
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)

        self.project.refresh_from_db()
        self.customer.refresh_from_db()
        task.refresh_from_db()
        self.assertEqual(task.status, FollowUpTask.Status.COMPLETED)
        self.assertEqual(self.project.progress, PROGRESS_STAGES[4])
        self.assertEqual(self.customer.progress, PROGRESS_STAGES[4])
        self.assertFalse(
            self.project.follow_up_tasks.filter(
                target_progress=PROGRESS_STAGES[6],
                assignee=self.business,
                status=FollowUpTask.Status.PENDING,
            ).exists(),
        )

        detail_response = self.client.get(
            reverse('customer-detail', args=[self.customer.id]),
        )
        self.assertEqual(detail_response.status_code, 200)
        completed_record = next(
            item for item in detail_response.json()['followUps']
            if item['id'] == task.id
        )
        self.assertEqual(completed_record['result'], '已完成现场工作并上传记录')
        self.assertTrue(completed_record['completedLabel'])
        self.assertEqual(completed_record['fromProgress'], PROGRESS_STAGES[4])

    def test_completed_result_can_be_edited_without_replaying_progress(self):
        task = self.project.follow_up_tasks.get(
            target_progress=PROGRESS_STAGES[5],
        )
        self.client.force_login(self.technical)
        url = reverse('follow-up-task-detail', args=[task.id])
        response = self.client.patch(
            url,
            data=json.dumps({
                'status': FollowUpTask.Status.COMPLETED,
                'result': '首次完成记录',
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        task.refresh_from_db()
        original_completed_at = task.completed_at

        response = self.client.patch(
            url,
            data=json.dumps({'result': '修正后的完整跟进结果'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)

        task.refresh_from_db()
        self.project.refresh_from_db()
        self.assertEqual(task.result, '修正后的完整跟进结果')
        self.assertEqual(task.completed_at, original_completed_at)
        self.assertEqual(self.project.progress, PROGRESS_STAGES[4])

    def test_project_owner_can_create_manual_task_without_advancing_stage(self):
        self.client.force_login(self.business)
        payload = {
            'projectId': self.project.id,
            'title': '临时补充一组材料试件',
            'role': FollowUpTask.Role.BUSINESS,
            'dueAt': (timezone.now() + timedelta(days=1)).isoformat(),
        }
        response = self.client.post(
            reverse('follow-up-tasks'),
            data=json.dumps(payload),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        created = response.json()['task']
        self.assertTrue(created['isManual'])
        self.assertEqual(created['fromProgress'], self.project.progress)
        self.assertEqual(created['targetProgress'], self.project.progress)

        task = FollowUpTask.objects.get(id=created['id'])
        response = self.client.patch(
            reverse('follow-up-task-detail', args=[task.id]),
            data=json.dumps({
                'status': FollowUpTask.Status.COMPLETED,
                'result': '临时任务已经处理完成',
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)

        task.refresh_from_db()
        self.project.refresh_from_db()
        self.customer.refresh_from_db()
        self.assertEqual(task.status, FollowUpTask.Status.COMPLETED)
        self.assertEqual(self.project.progress, PROGRESS_STAGES[4])
        self.assertEqual(self.customer.progress, PROGRESS_STAGES[4])

    def test_employee_cannot_create_manual_task_for_unassigned_project(self):
        self.client.force_login(self.other)
        response = self.client.post(
            reverse('follow-up-tasks'),
            data=json.dumps({
                'projectId': self.project.id,
                'title': '无权新增的项目任务',
                'role': FollowUpTask.Role.BUSINESS,
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 403)

    def test_project_owner_can_edit_task_fields_without_changing_progress(self):
        task = self.project.follow_up_tasks.get(
            target_progress=PROGRESS_STAGES[5],
        )
        due_at = timezone.now() + timedelta(days=2)
        self.client.force_login(self.business)

        response = self.client.patch(
            reverse('follow-up-task-detail', args=[task.id]),
            data=json.dumps({
                'title': '调整后的阶段跟进任务',
                'role': FollowUpTask.Role.BUSINESS,
                'dueAt': due_at.isoformat(),
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        task.refresh_from_db()
        self.project.refresh_from_db()
        self.customer.refresh_from_db()
        self.assertEqual(task.title, '调整后的阶段跟进任务')
        self.assertEqual(task.role, FollowUpTask.Role.BUSINESS)
        self.assertEqual(task.assignee, self.business)
        self.assertEqual(task.due_at, due_at)
        self.assertEqual(self.project.progress, PROGRESS_STAGES[4])
        self.assertEqual(self.customer.progress, PROGRESS_STAGES[4])
        self.assertTrue(response.json()['task']['canDelete'])

    def test_deleting_completed_record_does_not_roll_back_project(self):
        task = self.project.follow_up_tasks.get(
            target_progress=PROGRESS_STAGES[5],
        )
        self.client.force_login(self.technical)
        detail_url = reverse('follow-up-task-detail', args=[task.id])
        response = self.client.patch(
            detail_url,
            data=json.dumps({
                'status': FollowUpTask.Status.COMPLETED,
                'result': '阶段工作已经完成',
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)

        self.client.force_login(self.admin)
        response = self.client.delete(detail_url)

        self.assertEqual(response.status_code, 200)
        self.assertFalse(FollowUpTask.objects.filter(id=task.id).exists())
        self.project.refresh_from_db()
        self.customer.refresh_from_db()
        self.assertEqual(self.project.progress, PROGRESS_STAGES[4])
        self.assertEqual(self.customer.progress, PROGRESS_STAGES[4])


class PartnershipAgencyContractAndMapApiTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username='commercial-admin', password='123456',
        )
        self.agent = get_user_model().objects.create_user(
            username='regional-agent', password='123456',
        )
        EmployeeProfile.objects.create(
            user=self.agent, role=EmployeeProfile.Role.AGENT, nickname='浙江代理',
        )
        self.zhejiang = Customer.objects.create(
            name='浙江客户', phone='13800000041', province='浙江省', city='杭州市', district='西湖区', grade='A',
        )
        self.jiangsu = Customer.objects.create(
            name='江苏客户', phone='13800000042', province='江苏省', city='南京市', district='鼓楼区', grade='B',
        )
        self.primary_project = Project.objects.create(
            customer=self.zhejiang, name='浙江一期', progress='需求对接',
            province='浙江省', city='杭州市', district='西湖区',
        )
        self.client.force_login(self.admin)

    def authorization_payload(self, **overrides):
        payload = {
            'level': 'province',
            'province': '浙江省',
            'city': '',
            'district': '',
            'isExclusive': True,
            'productScope': '土凝岩材料全系列',
            'effectiveDate': timezone.localdate().isoformat(),
            'expiryDate': (timezone.localdate() + timedelta(days=60)).isoformat(),
            'agreementStatus': 'active',
            'agreementNumber': 'AUTH-001',
        }
        payload.update(overrides)
        return payload

    def test_selecting_partnership_identity_replaces_active_identity_and_keeps_history(self):
        for identity_type in ('customer', 'partner', 'technical_partner'):
            response = self.client.post(
                reverse('customer-partnerships', args=(self.zhejiang.id,)),
                data=json.dumps({'type': identity_type}),
                content_type='application/json',
            )
            self.assertIn(response.status_code, (200, 201))
        self.assertEqual(
            set(self.zhejiang.partnership_identities.filter(is_active=True).values_list('identity_type', flat=True)),
            {'technical_partner'},
        )
        self.assertEqual(self.zhejiang.partnership_identities.count(), 3)

    def test_exclusive_authorization_conflict_is_rejected_without_overwrite(self):
        created = self.client.post(
            reverse('customer-authorizations', args=(self.zhejiang.id,)),
            data=json.dumps(self.authorization_payload()),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 201)
        conflict = self.client.post(
            reverse('customer-authorizations', args=(self.jiangsu.id,)),
            data=json.dumps(self.authorization_payload()),
            content_type='application/json',
        )
        self.assertEqual(conflict.status_code, 409)
        self.assertIn('授权冲突', conflict.json()['error'])
        self.assertEqual(AgencyAuthorization.objects.count(), 1)

    def test_saved_agent_region_is_immediately_available_to_agent_map(self):
        created = self.client.post(
            reverse('customer-authorizations', args=(self.zhejiang.id,)),
            data=json.dumps(self.authorization_payload(
                level='city',
                city='杭州市',
                isExclusive=False,
            )),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 201)
        map_response = self.client.get(
            reverse('regional-business-map'),
            {'view': 'agents', 'province': '浙江省'},
        )
        self.assertEqual(map_response.status_code, 200)
        payload = map_response.json()
        self.assertEqual(payload['total'], 1)
        self.assertEqual(payload['regions'][0]['name'], '浙江省')
        self.assertEqual(payload['authorizations'][0]['city'], '杭州市')
        self.assertEqual(payload['authorizations'][0]['customerId'], self.zhejiang.id)

    def test_contract_has_independent_status_and_sixty_day_reminder(self):
        response = self.client.post(
            reverse('customer-contracts', args=(self.zhejiang.id,)),
            data=json.dumps({
                'title': '浙江年度供货合同',
                'projectId': self.primary_project.id,
                'status': 'active',
                'effectiveDate': timezone.localdate().isoformat(),
                'expiryDate': (timezone.localdate() + timedelta(days=45)).isoformat(),
                'amount': '120000.50',
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        contract = response.json()['contract']
        self.assertEqual(contract['status'], BusinessContract.Status.ACTIVE)
        self.assertEqual(contract['displayStatus'], BusinessContract.Status.EXPIRING)
        self.assertEqual(contract['expiryTone'], 'warning')
        self.assertIn('60天提醒', contract['reminderLabel'])
        self.assertTrue(self.zhejiang.partnership_identities.filter(
            identity_type=PartnershipIdentity.IdentityType.SIGNED_CUSTOMER,
        ).exists())

    def test_projects_keep_progress_updates_isolated(self):
        second = self.client.post(
            reverse('customer-projects', args=(self.zhejiang.id,)),
            data=json.dumps({'name': '浙江二期', 'progress': '技术验证'}),
            content_type='application/json',
        )
        self.assertEqual(second.status_code, 201)
        second_id = second.json()['project']['id']
        ProjectProgressUpdate.objects.create(
            project=self.primary_project,
            content='一期日常跟进',
            from_progress='需求对接',
            to_progress='技术验证',
            created_by=self.admin,
        )
        ProjectProgressUpdate.objects.create(
            project_id=second_id,
            content='二期日常跟进',
            from_progress='技术验证',
            to_progress='客户深度沟通',
            created_by=self.admin,
        )
        first_detail = self.client.get(
            reverse('customer-detail', args=(self.zhejiang.id,)),
            {'projectId': self.primary_project.id},
        ).json()
        second_detail = self.client.get(
            reverse('customer-detail', args=(self.zhejiang.id,)),
            {'projectId': second_id},
        ).json()
        self.assertEqual([item['content'] for item in first_detail['progressUpdates']], ['一期日常跟进'])
        self.assertEqual([item['content'] for item in second_detail['progressUpdates']], ['二期日常跟进'])

    def test_project_case_association_uses_exact_id_and_unlink_preserves_project(self):
        other_project = Project.objects.create(
            customer=self.jiangsu,
            name='江苏道路案例',
            progress='方案与报价',
        )
        linked = self.client.post(
            reverse('customer-project-associations', args=(self.zhejiang.id,)),
            data=json.dumps({'projectId': other_project.id}),
            content_type='application/json',
        )
        self.assertEqual(linked.status_code, 201)
        self.zhejiang.refresh_from_db()
        self.assertEqual(self.zhejiang.cooperation_status, Customer.CooperationStatus.COOPERATING)
        association = linked.json()['association']
        self.assertEqual(association['project']['id'], other_project.id)
        self.assertEqual(association['project']['customer']['id'], self.jiangsu.id)

        detail = self.client.get(
            reverse('customer-detail', args=(self.zhejiang.id,)),
            {'projectId': self.primary_project.id},
        ).json()
        self.assertIn(other_project.id, [item['project']['id'] for item in detail['associatedProjects']])

        unlinked = self.client.delete(
            reverse('customer-project-association-detail', args=(association['associationId'],)),
        )
        self.assertEqual(unlinked.status_code, 200)
        self.assertFalse(CustomerProjectAssociation.objects.filter(
            customer=self.zhejiang, project=other_project,
        ).exists())
        self.assertTrue(Project.objects.filter(id=other_project.id).exists())

    def test_project_completion_is_explicit_and_reversible(self):
        detail_url = reverse('project-detail', args=(self.primary_project.id,))

        completed = self.client.patch(
            detail_url,
            data=json.dumps({'isActive': False}),
            content_type='application/json',
        )
        self.assertEqual(completed.status_code, 200)
        self.assertFalse(completed.json()['project']['isActive'])
        self.primary_project.refresh_from_db()
        self.assertFalse(self.primary_project.is_active)

        restored = self.client.patch(
            detail_url,
            data=json.dumps({'isActive': True}),
            content_type='application/json',
        )
        self.assertEqual(restored.status_code, 200)
        self.assertTrue(restored.json()['project']['isActive'])
        self.primary_project.refresh_from_db()
        self.assertTrue(self.primary_project.is_active)

    def test_non_commercial_account_cannot_change_project_case_associations(self):
        authorization = AgencyAuthorization.objects.create(
            customer=self.zhejiang,
            level=AgencyAuthorization.Level.PROVINCE,
            province='浙江省',
            product_scope='全部产品',
            effective_date=timezone.localdate(),
            expiry_date=timezone.localdate() + timedelta(days=365),
            agreement_status=AgencyAuthorization.AgreementStatus.ACTIVE,
        )
        authorization.viewers.add(self.agent)
        self.client.force_login(self.agent)
        response = self.client.post(
            reverse('customer-project-associations', args=(self.zhejiang.id,)),
            data=json.dumps({'projectId': self.primary_project.id}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 403)

    def test_customer_edit_preserves_existing_primary_project_name(self):
        original_name = self.primary_project.name
        response = self.client.patch(
            reverse('customer-detail', args=(self.zhejiang.id,)),
            data=json.dumps({
                'name': self.zhejiang.name,
                'phone': self.zhejiang.phone,
                'source': self.zhejiang.source,
                'province': self.zhejiang.province,
                'city': self.zhejiang.city,
                'district': self.zhejiang.district,
                'grade': self.zhejiang.grade,
                'description': '客户资料保存回归验证',
                'progress': self.primary_project.progress,
                'plan': '更新后的方案说明',
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.primary_project.refresh_from_db()
        self.assertEqual(self.primary_project.name, original_name)
        self.assertEqual(self.primary_project.plan, '更新后的方案说明')

    def test_project_name_edit_targets_exact_project_and_deep_link_returns_it(self):
        second_project = Project.objects.create(
            customer=self.zhejiang,
            name='浙江二期原名',
            progress='技术验证',
            province='浙江省',
            city='杭州市',
            district='西湖区',
        )
        original_primary_name = self.primary_project.name

        response = self.client.patch(
            reverse('project-detail', args=(second_project.id,)),
            data=json.dumps({'name': '浙江二期更名'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.primary_project.refresh_from_db()
        second_project.refresh_from_db()
        self.assertEqual(self.primary_project.name, original_primary_name)
        self.assertEqual(second_project.name, '浙江二期更名')

        detail = self.client.get(
            reverse('customer-detail', args=(self.zhejiang.id,)),
            {'projectId': second_project.id},
        )
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.json()['selectedProjectId'], second_project.id)
        selected = next(
            project for project in detail.json()['projects']
            if project['id'] == second_project.id
        )
        self.assertEqual(selected['name'], '浙江二期更名')

    def test_project_name_cannot_be_cleared(self):
        original_name = self.primary_project.name
        response = self.client.patch(
            reverse('project-detail', args=(self.primary_project.id,)),
            data=json.dumps({'name': '   '}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()['error'], '项目名称不能为空')
        self.primary_project.refresh_from_db()
        self.assertEqual(self.primary_project.name, original_name)

    def test_agent_only_sees_customers_and_map_inside_linked_authorization(self):
        authorization = AgencyAuthorization.objects.create(
            customer=self.zhejiang,
            level=AgencyAuthorization.Level.PROVINCE,
            province='浙江省',
            product_scope='全部产品',
            effective_date=timezone.localdate(),
            expiry_date=timezone.localdate() + timedelta(days=365),
            agreement_status=AgencyAuthorization.AgreementStatus.ACTIVE,
        )
        authorization.viewers.add(self.agent)
        self.client.force_login(self.agent)
        customer_names = {
            item['name'] for item in self.client.get(reverse('customers-collection')).json()['customers']
        }
        self.assertIn('浙江客户', customer_names)
        self.assertNotIn('江苏客户', customer_names)
        self.assertTrue(all(
            item['province'] == '浙江省'
            for item in self.client.get(reverse('customers-collection')).json()['customers']
        ))
        map_response = self.client.get(
            reverse('regional-business-map'), {'view': 'customers'},
        )
        self.assertEqual(map_response.status_code, 200)
        self.assertEqual(map_response.json()['accessScope'], 'authorized_agent')
        self.assertEqual([item['name'] for item in map_response.json()['regions']], ['浙江省'])
        forbidden = self.client.get(reverse('customer-detail', args=(self.jiangsu.id,)))
        self.assertEqual(forbidden.status_code, 403)


class TomorrowItemApiTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='tomorrow-user',
            password='123456',
        )
        self.other = get_user_model().objects.create_user(
            username='tomorrow-other',
            password='123456',
        )
        self.client.force_login(self.user)

    def test_user_can_create_list_complete_and_delete_own_item(self):
        planned_date = (timezone.localdate() + timedelta(days=1)).isoformat()
        response = self.client.post(
            reverse('tomorrow-items'),
            data=json.dumps({
                'title': '整理明日客户回访名单',
                'note': '优先处理A级客户',
                'plannedDate': planned_date,
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        item_id = response.json()['item']['id']

        response = self.client.get(
            reverse('tomorrow-items'),
            {'date': planned_date},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()['items']), 1)
        self.assertEqual(response.json()['items'][0]['title'], '整理明日客户回访名单')

        response = self.client.patch(
            reverse('tomorrow-item-detail', args=[item_id]),
            data=json.dumps({'isCompleted': True}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['item']['isCompleted'])

        response = self.client.delete(
            reverse('tomorrow-item-detail', args=[item_id]),
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(TomorrowItem.objects.filter(id=item_id).exists())

    def test_items_are_private_to_the_logged_in_account(self):
        private_item = TomorrowItem.objects.create(
            user=self.other,
            title='其他员工的明日事项',
        )
        response = self.client.get(reverse('tomorrow-items'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['items'], [])

        response = self.client.patch(
            reverse('tomorrow-item-detail', args=[private_item.id]),
            data=json.dumps({'isCompleted': True}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 404)

    def test_item_title_is_required(self):
        response = self.client.post(
            reverse('tomorrow-items'),
            data=json.dumps({'title': ''}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)
